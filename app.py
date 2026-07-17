import streamlit as st
import pandas as pd
from pathlib import Path
from datetime import date, datetime
import uuid

BASE_DIR = Path(__file__).parent
DADOS_DIR = BASE_DIR / "dados"
ASSETS_DIR = BASE_DIR / "assets"

ELO_AZUL = ASSETS_DIR / "elo_azul.svg"

st.set_page_config(
    page_title="Precificação da DOA | Facto",
    page_icon=str(ELO_AZUL),
    layout="wide",
    initial_sidebar_state="expanded"
)

# =========================
# CORES FACTO
# =========================

VERMELHO = "#D34022"
AZUL = "#163246"
OFF_WHITE = "#EAE8DD"
BRANCO = "#FFFFFF"

# =========================
# CAMINHOS
# =========================

BASE_DIR = Path(__file__).parent
DADOS_DIR = BASE_DIR / "dados"
ASSETS_DIR = BASE_DIR / "assets"

DADOS_DIR.mkdir(exist_ok=True)
ASSETS_DIR.mkdir(exist_ok=True)

ARQ_PROJETOS = DADOS_DIR / "projetos.xlsx"
ARQ_PARF = DADOS_DIR / "parf_itens.xlsx"
ARQ_ICP_GIRO = DADOS_DIR / "icp_giro.xlsx"
ARQ_DOA = DADOS_DIR / "doa_resultados.xlsx"
ARQ_AUXILIARES = DADOS_DIR / "auxiliares.xlsx"
ARQ_CUSTOS_INDICADORES = DADOS_DIR / "custos_indicadores.xlsx"
ARQ_MODELO_PROPOSTA = DADOS_DIR / "modelo_proposta.docx"
SAIDAS_DIR = BASE_DIR / "saidas"
SAIDAS_DIR.mkdir(exist_ok=True)

LOGO_FACTO = ASSETS_DIR / "logo_facto.svg"

# =========================
# CSS
# =========================

st.markdown(
    f"""
    <style>

    .stApp {{
        background-color: white;
        color: {AZUL};
    }}

    .block-container {{
        background-color: white;
    }}

    p, span, label, div {{
        color: {AZUL};
    }}

    h1, h2, h3 {{
        color: {AZUL};
    }}

    section[data-testid="stSidebar"] {{
        background-color: {AZUL};
    }}

    section[data-testid="stSidebar"] label,
    section[data-testid="stSidebar"] p,
    section[data-testid="stSidebar"] span,
    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3 {{
        color: white !important;
    }}

    section[data-testid="stSidebar"] input,
    section[data-testid="stSidebar"] textarea,
    section[data-testid="stSidebar"] select,
    section[data-testid="stSidebar"] div[data-baseweb="select"] *,
    section[data-testid="stSidebar"] div[data-baseweb="input"] * {{
        color: {AZUL} !important;
    }}

    section[data-testid="stSidebar"] button {{
        color: {AZUL} !important;
        background-color: {OFF_WHITE} !important;
        border: 1px solid {OFF_WHITE} !important;
        font-weight: 700 !important;
        border-radius: 10px !important;
    }}

    section[data-testid="stSidebar"] button:hover {{
        color: white !important;
        background-color: {VERMELHO} !important;
        border: 1px solid {VERMELHO} !important;
    }}

    .facto-title {{
        color: white;
        font-size: 34px;
        font-weight: 800;
        margin-bottom: 4px;
    }}

    .facto-subtitle {{
        color: {OFF_WHITE};
        font-size: 16px;
    }}

    .facto-badge {{
        display: inline-block;
        background-color: {VERMELHO};
        color: white;
        padding: 6px 14px;
        border-radius: 999px;
        font-size: 12px;
        font-weight: 700;
        margin-bottom: 10px;
    }}

    div[data-testid="stForm"] {{
        background-color: white;
        padding: 26px;
        border-radius: 22px;
        border: 1px solid #ddd;
        box-shadow: 0 6px 18px rgba(22, 50, 70, 0.10);
    }}

    .metric-card {{
        background: white;
        padding: 22px;
        border-radius: 18px;
        border: 1px solid #ddd;
        box-shadow: 0 4px 14px rgba(22, 50, 70, 0.08);
    }}

    .metric-label {{
        color: #5f6f7a;
        font-size: 13px;
        font-weight: 700;
        text-transform: uppercase;
    }}

    .metric-value {{
        color: {AZUL};
        font-size: 26px;
        font-weight: 800;
    }}

    .facto-callout {{
        border-left: 6px solid #D34022;
        background: #fff5f2;
        padding: 16px 18px;
        border-radius: 12px;
        margin-top: 12px;
        color: #163246 !important;
        font-family: inherit !important;
        font-size: 15px;
        line-height: 1.5;
    }}

    .facto-callout strong {{
        color: #D34022 !important;
        font-weight: 800;
    }}

    .facto-info-card {{
        background: #F7F9FB;
        border: 1px solid #DDE4EA;
        border-radius: 14px;
        padding: 14px 16px;
        margin-top: 8px;
        color: #163246 !important;
        font-family: inherit !important;
        font-size: 15px;
        line-height: 1.5;
    }}

    .facto-info-card strong {{
        color: #163246 !important;
        font-weight: 800;
    }}

    </style>
    """,
    unsafe_allow_html=True
)

# =========================
# FUNÇÕES BASE
# =========================

def moeda(valor):
    try:
        return f"R$ {float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except Exception:
        return "R$ 0,00"


def carregar_excel(path: Path, colunas: list) -> pd.DataFrame:
    if path.exists():
        return pd.read_excel(path)
    return pd.DataFrame(columns=colunas)


def salvar_excel(df: pd.DataFrame, path: Path):
    df.to_excel(path, index=False)


def gerar_id():
    return str(uuid.uuid4())[:8]


def agora():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def colunas_projetos():
    return [
        "id_projeto", "versao", "status",
        "nome_projeto", "coordenador", "email_coordenador",
        "instituicao_executora", "campus", "financiador",
        "tipo_financiador", "tipo_instrumento",
        "data_inicio", "data_fim", "prazo_meses",
        "valor_aprovado", "objeto_resumido", "area_tematica",
        "possui_bolsas", "possui_compras", "possui_concurso_ps",
        "possui_clt_rpa_pj", "possui_importacao",
        "responsavel_facto", "data_criacao", "data_atualizacao"
    ]


def colunas_parf():
    return [
        "id_item",
        "id_projeto", "versao", "grupo", "item",
        "modalidade", "categoria", "descricao",
        "quantidade", "meses", "valor_unitario",
        "percentual_adicional", "valor_adicional",
        "modalidade_contratacao", "total"
    ]


def colunas_icp_giro():
    return [
        "id_projeto", "versao", "tipo", "indicador",
        "nota", "justificativa"
    ]


def colunas_doa():
    return [
        "id_projeto", "versao", "componente", "setor", "atividade",
        "quantidade_operacional", "tempo_base_h", "vigencia_meses", "horas_estimadas",
        "custo_hora_pessoal", "custo_hora_adm", "custo_hora_total",
        "fator_icp", "fator_giro", "fator_complexidade",
        "valor_calculado", "valor_ajustado",
        "fator_limitador", "data_calculo"
    ]


def carregar_projetos():
    return carregar_excel(ARQ_PROJETOS, colunas_projetos())


def carregar_parf():
    df = carregar_excel(ARQ_PARF, colunas_parf())

    # Compatibilidade com arquivos antigos sem id_item
    if "id_item" not in df.columns:
        df.insert(0, "id_item", [gerar_id() for _ in range(len(df))])
        if not df.empty:
            salvar_excel(df, ARQ_PARF)

    # Garante todas as colunas esperadas
    for coluna in colunas_parf():
        if coluna not in df.columns:
            df[coluna] = None

    return df[colunas_parf()]


def carregar_icp_giro():
    return carregar_excel(ARQ_ICP_GIRO, colunas_icp_giro())


def carregar_doa():
    return carregar_excel(ARQ_DOA, colunas_doa())


def calcular_prazo_meses(inicio, fim):
    if not inicio or not fim:
        return 0
    return max(1, (fim.year - inicio.year) * 12 + (fim.month - inicio.month) + 1)


def parse_data_streamlit(valor, padrao=None):
    """Converte datas salvas em texto/Excel para date, preservando cadastro ao reabrir proposta."""
    if padrao is None:
        padrao = date.today()
    if valor is None or valor == "":
        return padrao
    try:
        if pd.isna(valor):
            return padrao
    except Exception:
        pass
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        convertido = pd.to_datetime(valor, errors="coerce")
        if pd.isna(convertido):
            return padrao
        return convertido.date()
    except Exception:
        return padrao


def selectbox_obrigatorio(label, opcoes, key=None, texto_vazio="Selecione..."):
    """Lista suspensa pesquisável, sem dependência externa.

    O usuário digita parte do texto e a lista é filtrada dinamicamente.
    Esta função deve ser usada fora de st.form para que a busca atualize
    imediatamente a cada digitação.
    """
    opcoes_limpas = limpar_lista(opcoes)
    chave_base = key or nome_seguro_arquivo(label)

    termo = st.text_input(
        f"Buscar em {label.lower()}",
        value="",
        placeholder="Digite parte do texto para filtrar...",
        key=f"{chave_base}_busca"
    ).strip()

    if termo:
        termo_norm = termo.casefold()
        opcoes_filtradas = [
            opcao for opcao in opcoes_limpas
            if termo_norm in str(opcao).casefold()
        ]
    else:
        opcoes_filtradas = opcoes_limpas

    if termo and not opcoes_filtradas:
        st.caption("Nenhuma opção encontrada para a busca informada.")

    # A chave da seleção varia conforme o termo, evitando que o Streamlit
    # retenha uma opção que deixou de existir após a filtragem.
    termo_chave = nome_seguro_arquivo(termo) if termo else "todos"
    valor = st.selectbox(
        label,
        [texto_vazio] + opcoes_filtradas,
        index=0,
        key=f"{chave_base}_selecao_{termo_chave}"
    )
    return "" if valor == texto_vazio else valor


def obter_vigencia_projeto_atual():
    projeto = obter_projeto_atual() or {}
    try:
        return max(1.0, float(projeto.get("prazo_meses", 1) or 1))
    except Exception:
        return 1.0


def calcular_operacoes_bancarias_estimadas(df_parf_atual):
    if df_parf_atual is None or df_parf_atual.empty:
        return 0.0
    total_ops = 0.0
    for _, row in df_parf_atual.iterrows():
        grupo = str(row.get("grupo", ""))
        if grupo == "Tarifas bancárias":
            continue
        qtd = pd.to_numeric(row.get("quantidade", 0), errors="coerce")
        meses = pd.to_numeric(row.get("meses", 1), errors="coerce")
        qtd = 0.0 if pd.isna(qtd) else float(qtd)
        meses = 1.0 if pd.isna(meses) or float(meses) <= 0 else float(meses)
        if grupo in ["Bolsas", "Celetistas"]:
            total_ops += qtd * meses
        elif grupo == "RPA":
            total_ops += qtd
        elif grupo in ["Material de consumo", "Material permanente", "Serviço PJ", "Importação", "Diárias", "Passagens", "Ressarcimento", "Contrapartida", "Prospecção"]:
            total_ops += max(1.0, qtd)
    return total_ops


def calcular_tarifas_bancarias_estimadas(df_parf_atual, vigencia_meses):
    operacoes = calcular_operacoes_bancarias_estimadas(df_parf_atual)
    tarifa_operacao = 2.06
    tarifa_manutencao_mensal = 100.00
    valor_operacoes = operacoes * tarifa_operacao
    valor_manutencao = max(1.0, float(vigencia_meses or 1)) * tarifa_manutencao_mensal
    total = valor_operacoes + valor_manutencao
    return operacoes, valor_operacoes, valor_manutencao, total


def carregar_auxiliares():
    aux = {}

    if not ARQ_AUXILIARES.exists():
        return aux

    try:
        xls = pd.ExcelFile(ARQ_AUXILIARES)
        for aba in xls.sheet_names:
            df = pd.read_excel(ARQ_AUXILIARES, sheet_name=aba)
            chave_original = str(aba)
            chave_limpa = chave_original.strip()
            aux[chave_original] = df
            aux[chave_limpa] = df
    except Exception:
        pass

    return aux


def grupos_parf():
    return [
        "Bolsas",
        "Celetistas",
        "RPA",
        "Material de consumo",
        "Material permanente",
        "Serviço PJ",
        "Importação",
        "Diárias",
        "Passagens",
        "Tarifas bancárias",
        "Ressarcimento",
        "Contrapartida",
        "Prospecção"
    ]


def grupos_contratacao():
    return ["Material de consumo", "Material permanente", "Serviço PJ"]


def limpar_lista(valores):
    proibidos = [
        "nan", "none", "subelementos de despesa", "equipamentos e materiais permanentes",
        "material de consumo", "serviços de terceiros", "tipo de despesa"
    ]
    lista = []
    for valor in valores:
        if pd.isna(valor):
            continue
        texto = str(valor).strip()
        if not texto:
            continue
        if texto.lower() in proibidos:
            continue
        if texto not in lista:
            lista.append(texto)
    return lista



def _nome_aba_por_prefixo(prefixo):
    """Localiza uma aba no arquivo auxiliares.xlsx por prefixo, ignorando espaços."""
    if not ARQ_AUXILIARES.exists():
        return None
    try:
        xls = pd.ExcelFile(ARQ_AUXILIARES)
        alvo = str(prefixo).strip().upper()
        for aba in xls.sheet_names:
            if str(aba).strip().upper().startswith(alvo):
                return aba
    except Exception:
        return None
    return None


def limpar_lista(valores):
    lista = []
    for v in valores:
        if pd.isna(v):
            continue
        texto = str(v).strip()
        if not texto or texto.lower() in ["nan", "none"]:
            continue
        if texto.upper() in [
            "MODALIDADE",
            "CATEGORIA",
            "SUBELEMENTOS DE DESPESA",
            "EQUIPAMENTOS E MATERIAIS PERMANENTES",
            "MATERIAL DE CONSUMO",
            "SERVIÇOS DE TERCEIROS",
            "SERVICOS DE TERCEIROS",
            "TIPO DE DESPESA",
            "TIPO DE PASSAGENS",
        ]:
            continue
        lista.append(texto)
    return sorted(list(dict.fromkeys(lista)))


def _normalizar_coluna_nome(coluna):
    texto = str(coluna or "").strip().lower()
    troca = {
        "á":"a", "à":"a", "ã":"a", "â":"a",
        "é":"e", "ê":"e",
        "í":"i",
        "ó":"o", "ô":"o", "õ":"o",
        "ú":"u",
        "ç":"c",
    }
    for origem, destino in troca.items():
        texto = texto.replace(origem, destino)
    texto = texto.replace(" ", "_").replace("-", "_")
    while "__" in texto:
        texto = texto.replace("__", "_")
    return texto.strip("_")


def _ler_aba_normalizada(nome_aba):
    if not ARQ_AUXILIARES.exists():
        return pd.DataFrame()
    try:
        xls = pd.ExcelFile(ARQ_AUXILIARES)
        mapa = {str(aba).strip().casefold(): aba for aba in xls.sheet_names}
        aba_real = mapa.get(str(nome_aba).strip().casefold())
        if not aba_real:
            return pd.DataFrame()
        df = pd.read_excel(ARQ_AUXILIARES, sheet_name=aba_real)
        df.columns = [_normalizar_coluna_nome(c) for c in df.columns]
        return df
    except Exception:
        return pd.DataFrame()


def obter_opcoes_bolsas(aux=None):
    """Lê preferencialmente a aba normalizada aux_bolsas.

    Esperado:
    - modalidade
    - categoria
    - valor_referencia (opcional)

    Mantém fallback para a aba antiga CLASSIFICAÇÃO DE BOLSAS.
    """
    df_norm = _ler_aba_normalizada("aux_bolsas")
    if not df_norm.empty and {"modalidade", "categoria"}.issubset(set(df_norm.columns)):
        col_valor = None
        for possivel in ["valor_referencia", "valor", "valor_unitario", "valor_bolsa"]:
            if possivel in df_norm.columns:
                col_valor = possivel
                break

        dados = pd.DataFrame({
            "Modalidade": df_norm["modalidade"],
            "Categoria": df_norm["categoria"],
            "Valor": df_norm[col_valor] if col_valor else 0,
        })

        dados["Modalidade"] = dados["Modalidade"].astype(str).str.strip()
        dados["Categoria"] = dados["Categoria"].astype(str).str.strip()
        dados["Valor"] = pd.to_numeric(dados["Valor"], errors="coerce").fillna(0)
        dados = dados[
            ~dados["Modalidade"].str.lower().isin(["nan", "none", ""])
            & ~dados["Categoria"].str.lower().isin(["nan", "none", ""])
        ].copy()
        return dados.drop_duplicates().reset_index(drop=True)

    # Fallback: estrutura visual antiga.
    aba = _nome_aba_por_prefixo("CLASSIFICAÇÃO DE BOLSAS")
    if not aba:
        return pd.DataFrame(columns=["Modalidade", "Categoria", "Valor"])

    try:
        raw = pd.read_excel(ARQ_AUXILIARES, sheet_name=aba, header=None)
    except Exception:
        return pd.DataFrame(columns=["Modalidade", "Categoria", "Valor"])

    if raw.empty or raw.shape[1] < 3:
        return pd.DataFrame(columns=["Modalidade", "Categoria", "Valor"])

    header = raw.iloc[0].astype(str).str.replace("\xa0", " ", regex=False).str.strip().tolist()
    col_modalidade = 0
    col_categoria = 2
    col_valor = None

    for i, h in enumerate(header):
        h_norm = h.lower()
        if "valor" in h_norm:
            col_valor = i
            break

    if col_valor is None and raw.shape[1] >= 9:
        col_valor = 8

    dados = pd.DataFrame({
        "Modalidade": raw.iloc[1:, col_modalidade],
        "Categoria": raw.iloc[1:, col_categoria],
        "Valor": raw.iloc[1:, col_valor] if col_valor is not None and col_valor < raw.shape[1] else 0
    })

    dados["Modalidade"] = dados["Modalidade"].astype(str).str.strip()
    dados["Categoria"] = dados["Categoria"].astype(str).str.strip()
    dados["Valor"] = pd.to_numeric(dados["Valor"], errors="coerce").fillna(0)
    dados = dados[
        ~dados["Modalidade"].str.lower().isin(["nan", "none", ""])
        & ~dados["Categoria"].str.lower().isin(["nan", "none", ""])
    ].copy()
    return dados.drop_duplicates().reset_index(drop=True)


def obter_valor_bolsa(df_bolsas, modalidade, categoria):
    if df_bolsas.empty:
        return 0.0

    df = df_bolsas.copy()

    df["Modalidade"] = df["Modalidade"].astype(str).str.strip()
    df["Categoria"] = df["Categoria"].astype(str).str.strip()

    filtro = (
        df["Modalidade"].str.casefold() == str(modalidade).strip().casefold()
    ) & (
        df["Categoria"].str.casefold() == str(categoria).strip().casefold()
    )

    resultado = df.loc[filtro, "Valor"]

    if resultado.empty:
        return 0.0

    return float(resultado.iloc[0] or 0)


def obter_lista_despesas(aux, grupo):
    """Lê preferencialmente a aba normalizada aux_despesas.

    Esperado:
    - grupo_parf
    - categoria

    Mantém fallback para a aba antiga CLASSIFICAÇÃO DE DESPESAS.
    """
    df_norm = _ler_aba_normalizada("aux_despesas")
    if not df_norm.empty and {"grupo_parf", "categoria"}.issubset(set(df_norm.columns)):
        temp = df_norm.copy()
        temp["grupo_norm"] = temp["grupo_parf"].astype(str).str.strip().str.casefold()
        grupo_norm = str(grupo).strip().casefold()
        valores = temp.loc[temp["grupo_norm"] == grupo_norm, "categoria"].tolist()
        return limpar_lista(valores)

    # Fallback: estrutura visual antiga.
    aba = _nome_aba_por_prefixo("CLASSIFICAÇÃO DE DESPESAS")
    if not aba:
        return []

    try:
        raw = pd.read_excel(ARQ_AUXILIARES, sheet_name=aba, header=None)
    except Exception:
        return []

    if raw.empty:
        return []

    raw = raw.dropna(axis=1, how="all")
    if raw.empty:
        return []

    linha_header = None
    for i in range(raw.shape[0]):
        linha = " | ".join([str(x).strip().upper() for x in raw.iloc[i].tolist() if pd.notna(x)])
        if "EQUIPAMENTOS" in linha and "MATERIAL DE CONSUMO" in linha:
            linha_header = i
            break

    if linha_header is None:
        linha_header = 0

    mapa_posicao = {
        "Material permanente": 0,
        "Material de consumo": 1,
        "Serviço PJ": 2,
    }

    pos = mapa_posicao.get(grupo)
    if pos is None or pos >= raw.shape[1]:
        return []

    valores = raw.iloc[linha_header + 1:, pos].tolist()
    return limpar_lista(valores)


def obter_lista_diarias(aux):
    """Lê preferencialmente a aba normalizada aux_diarias.

    Esperado:
    - tipo_diaria
    - valor_unitario
    """
    df_norm = _ler_aba_normalizada("aux_diarias")
    if not df_norm.empty and "tipo_diaria" in df_norm.columns:
        col_valor = "valor_unitario" if "valor_unitario" in df_norm.columns else ("valor" if "valor" in df_norm.columns else None)
        dados = pd.DataFrame({
            "tipo": df_norm["tipo_diaria"],
            "valor": df_norm[col_valor] if col_valor else 0,
        })
        dados["tipo"] = dados["tipo"].astype(str).str.strip()
        dados["valor"] = pd.to_numeric(dados["valor"], errors="coerce").fillna(0)
        dados = dados[~dados["tipo"].str.lower().isin(["nan", "none", ""])].copy()
        return dados.drop_duplicates().reset_index(drop=True)

    # Fallback: estrutura visual antiga.
    aba = _nome_aba_por_prefixo("CLASSIFICAÇÃO E VALORES DE DIÁR")
    if not aba:
        return pd.DataFrame(columns=["tipo", "valor"])

    try:
        raw = pd.read_excel(ARQ_AUXILIARES, sheet_name=aba, header=None)
    except Exception:
        return pd.DataFrame(columns=["tipo", "valor"])

    if raw.empty:
        return pd.DataFrame(columns=["tipo", "valor"])

    raw = raw.dropna(axis=1, how="all")
    if raw.shape[1] < 2:
        return pd.DataFrame(columns=["tipo", "valor"])

    linha_header = None
    for i in range(raw.shape[0]):
        vals = [str(x).strip().upper() for x in raw.iloc[i].tolist() if pd.notna(x)]
        if any("TIPO" in v for v in vals):
            linha_header = i
            break

    if linha_header is None:
        linha_header = 0

    linhas = []
    for i in range(linha_header + 1, raw.shape[0]):
        tipo = raw.iat[i, 0]
        valor = raw.iat[i, 1]
        valor_num = pd.to_numeric(valor, errors="coerce")
        if pd.isna(tipo) or pd.isna(valor_num):
            continue
        tipo_txt = str(tipo).strip()
        if not tipo_txt or tipo_txt.lower() in ["nan", "none"]:
            continue
        if "COTAÇÃO" in tipo_txt.upper() or "DÓLAR" in tipo_txt.upper():
            continue
        linhas.append({"tipo": tipo_txt, "valor": float(valor_num)})

    return pd.DataFrame(linhas).drop_duplicates().reset_index(drop=True)


def obter_lista_passagens(aux):
    """Lê preferencialmente a aba normalizada aux_passagens.

    Esperado:
    - tipo_passagem
    """
    df_norm = _ler_aba_normalizada("aux_passagens")
    if not df_norm.empty and "tipo_passagem" in df_norm.columns:
        return limpar_lista(df_norm["tipo_passagem"].tolist())

    aba = _nome_aba_por_prefixo("CLASSIFICAÇÃO PASSAGENS")
    if not aba:
        return []

    try:
        raw = pd.read_excel(ARQ_AUXILIARES, sheet_name=aba, header=None)
    except Exception:
        return []

    raw = raw.dropna(axis=1, how="all")
    if raw.empty:
        return []

    valores = raw.iloc[:, 0].tolist()
    return limpar_lista(valores)


def salvar_ou_atualizar_projeto(dados):
    df = carregar_projetos()

    # Garante que todas as colunas aceitem texto, número e vazio
    for col in df.columns:
        df[col] = df[col].astype("object")

    # Padroniza datas como texto
    for campo in ["data_inicio", "data_fim", "data_criacao", "data_atualizacao"]:
        if campo in dados and dados[campo] is not None:
            dados[campo] = str(dados[campo])

    # Substitui NaN por vazio apenas nos dados recebidos
    for k, v in list(dados.items()):
        if pd.isna(v):
            dados[k] = ""

    if not df.empty:
        filtro = (
            (df["id_projeto"].astype(str) == str(dados["id_projeto"])) &
            (df["versao"].astype(str) == str(dados["versao"]))
        )
    else:
        filtro = pd.Series([], dtype=bool)

    if filtro.any():
        idx = df.index[filtro][0]

        for k, v in dados.items():
            if k in df.columns:
                df.at[idx, k] = v
    else:
        df = pd.concat([df, pd.DataFrame([dados])], ignore_index=True)

    salvar_excel(df, ARQ_PROJETOS)


def obter_projeto_atual():
    id_projeto = st.session_state.get("id_projeto")
    versao = st.session_state.get("versao")

    if not id_projeto or not versao:
        return None

    df = carregar_projetos()

    if df.empty:
        return None

    filtro = (
        (df["id_projeto"].astype(str) == str(id_projeto)) &
        (df["versao"].astype(int) == int(versao))
    )

    if filtro.any():
        return df.loc[filtro].iloc[0].to_dict()

    return None


def mostrar_card_proposta():
    id_projeto = st.session_state.get("id_projeto")
    versao = st.session_state.get("versao")

    if not id_projeto or not versao:
        return

    projeto = obter_projeto_atual() or {}

    nome = projeto.get("nome_projeto", "Proposta ainda sem nome")
    status = projeto.get("status", st.session_state.get("status", "Rascunho"))
    valor = projeto.get("valor_aprovado", 0)

    with st.container(border=True):
        st.caption("PROPOSTA EM EDIÇÃO")
        st.markdown(f"### {nome}")
        col1, col2, col3, col4 = st.columns(4)

        with col1:
            st.markdown(f"**ID**  \n{id_projeto}")
        with col2:
            st.markdown(f"**Versão**  \n{versao}")
        with col3:
            st.markdown(f"**Status**  \n{status}")
        with col4:
            st.markdown(f"**Valor previsto**  \n{moeda(valor)}")


def atualizar_item_parf(id_item, dados_atualizados):
    df = carregar_parf()

    if df.empty or "id_item" not in df.columns:
        return False

    filtro = df["id_item"].astype(str) == str(id_item)

    if not filtro.any():
        return False

    idx = df.index[filtro][0]

    for k, v in dados_atualizados.items():
        if k in df.columns:
            df.at[idx, k] = v

    salvar_excel(df, ARQ_PARF)
    return True


def excluir_item_parf(id_item):
    df = carregar_parf()

    if df.empty or "id_item" not in df.columns:
        return False

    filtro = df["id_item"].astype(str) == str(id_item)

    if not filtro.any():
        return False

    df = df.loc[~filtro].copy()
    salvar_excel(df, ARQ_PARF)
    return True


def criar_nova_versao(id_projeto, versao_origem):
    nova_versao = int(versao_origem) + 1

    dfp = carregar_projetos()

    origem = dfp[
        (dfp["id_projeto"].astype(str) == str(id_projeto)) &
        (dfp["versao"].astype(int) == int(versao_origem))
    ]

    if origem.empty:
        return None

    novo = origem.iloc[0].copy()
    novo["versao"] = nova_versao
    novo["status"] = "Renegociada"
    novo["data_criacao"] = agora()
    novo["data_atualizacao"] = agora()

    dfp = pd.concat([dfp, pd.DataFrame([novo])], ignore_index=True)
    salvar_excel(dfp, ARQ_PROJETOS)

    for arquivo, colunas in [
        (ARQ_PARF, colunas_parf()),
        (ARQ_ICP_GIRO, colunas_icp_giro()),
        (ARQ_DOA, colunas_doa())
    ]:
        df = carregar_excel(arquivo, colunas)

        if df.empty:
            continue

        dados_origem = df[
            (df["id_projeto"].astype(str) == str(id_projeto)) &
            (df["versao"].astype(int) == int(versao_origem))
        ].copy()

        if not dados_origem.empty:
            dados_origem["versao"] = nova_versao
            df = pd.concat([df, dados_origem], ignore_index=True)
            salvar_excel(df, arquivo)

    return nova_versao



def excluir_proposta_versao(id_projeto, versao):
    """Exclui uma versão da proposta e todos os registros vinculados.

    A exclusão abrange cadastro, PARF, ICP/Giro, DOA e arquivos DOCX gerados
    para a mesma combinação de ID e versão.
    """
    id_txt = str(id_projeto)
    versao_int = int(versao)
    removidos = {}

    estruturas = [
        (ARQ_PROJETOS, colunas_projetos(), "projetos"),
        (ARQ_PARF, colunas_parf(), "parf"),
        (ARQ_ICP_GIRO, colunas_icp_giro(), "icp_giro"),
        (ARQ_DOA, colunas_doa(), "doa"),
    ]

    for arquivo, colunas, nome in estruturas:
        df = carregar_excel(arquivo, colunas)
        if df.empty or "id_projeto" not in df.columns or "versao" not in df.columns:
            removidos[nome] = 0
            continue

        versoes = pd.to_numeric(df["versao"], errors="coerce")
        filtro = (df["id_projeto"].astype(str) == id_txt) & (versoes == versao_int)
        removidos[nome] = int(filtro.sum())
        df_final = df.loc[~filtro].copy()

        for coluna in colunas:
            if coluna not in df_final.columns:
                df_final[coluna] = None
        salvar_excel(df_final[colunas], arquivo)

    removidos["documentos"] = 0
    if SAIDAS_DIR.exists():
        padrao = f"*_v{versao_int}_{id_txt}.docx"
        for arquivo_docx in SAIDAS_DIR.glob(padrao):
            try:
                arquivo_docx.unlink()
                removidos["documentos"] += 1
            except OSError:
                pass

    return removidos


# =========================
# FUNÇÕES DOA
# =========================

def carregar_custos_indicadores():
    """Carrega custos_indicadores.xlsx, que deve estar na pasta dados."""
    bases = {}

    if not ARQ_CUSTOS_INDICADORES.exists():
        return bases

    try:
        xls = pd.ExcelFile(ARQ_CUSTOS_INDICADORES)
        for aba in xls.sheet_names:
            df = pd.read_excel(ARQ_CUSTOS_INDICADORES, sheet_name=aba)
            df.columns = df.columns.astype(str).str.strip()
            bases[aba.strip()] = df
    except Exception:
        pass

    return bases


def normalizar_texto(valor):
    return str(valor or "").strip().lower()


def buscar_atividade(tempo_df, atividade):
    """Busca uma atividade na aba tempo_pessoas pelo nome exato; se não achar, tenta por contém."""
    if tempo_df.empty or "atividade" not in tempo_df.columns:
        return None

    alvo = normalizar_texto(atividade)
    temp = tempo_df.copy()
    temp["atividade_norm"] = temp["atividade"].apply(normalizar_texto)

    exata = temp[temp["atividade_norm"] == alvo]
    if not exata.empty:
        return exata.iloc[0].to_dict()

    contem = temp[temp["atividade_norm"].str.contains(alvo, na=False, regex=False)]
    if not contem.empty:
        return contem.iloc[0].to_dict()

    return None


def componente_por_setor(setor):
    mapa = {
        "Compras": "Gerenciamento de compras e contratações",
        "SelectID": "Gerenciamento de Bolsas/Auxílios e/ou PS e Concursos",
        "Projetos": "Gerenciamento de Projetos",
        "Financeiro": "Gerenciamento administrativo e financeiro",
        "Gente e Gestão": "Gestão institucional e governança",
    }
    return mapa.get(str(setor), str(setor))


def fator_governanca_por_valor(valor_projeto):
    """
    Fator institucional automático para representar governança, exposição,
    controles e risco associados ao porte financeiro do projeto.
    Não exige preenchimento pelo usuário.
    """
    try:
        valor = max(0.0, float(valor_projeto or 0))
    except Exception:
        valor = 0.0

    if valor <= 500_000:
        return 1.00
    if valor <= 2_000_000:
        return 1.10
    if valor <= 5_000_000:
        return 1.20
    if valor <= 10_000_000:
        return 1.30
    if valor <= 20_000_000:
        return 1.40
    return 1.50


def calcular_ajustes_metodologicos_doa(total_calculado, valor_projeto, prazo_meses):
    """
    Consolida a metodologia institucional da DOA:

    1. esforço operacional apurado pelas atividades;
    2. sustentação temporal amortecida: +25% por ano adicional;
    3. fator automático de governança conforme o porte financeiro;
    4. piso institucional de 5% do valor aprovado/previsto;
    5. teto legal de 15% do valor aprovado/previsto.

    Todos os parâmetros são automáticos e não aparecem como campos de entrada.
    """
    try:
        total_base = max(0.0, float(total_calculado or 0))
    except Exception:
        total_base = 0.0

    try:
        valor_base = max(0.0, float(valor_projeto or 0))
    except Exception:
        valor_base = 0.0

    try:
        meses = max(1.0, float(prazo_meses or 1))
    except Exception:
        meses = 1.0

    anos = max(1.0, meses / 12.0)
    fator_temporal = 1.0 + (0.25 * max(0.0, anos - 1.0))
    fator_governanca = fator_governanca_por_valor(valor_base)

    doa_metodologica = total_base * fator_temporal * fator_governanca
    piso_institucional = valor_base * 0.05
    limite_legal = valor_base * 0.15

    doa_antes_teto = max(doa_metodologica, piso_institucional)
    doa_final = min(doa_antes_teto, limite_legal) if valor_base > 0 else doa_antes_teto

    fator_ajuste_final = (doa_final / total_base) if total_base > 0 else 0.0
    percentual_final = (doa_final / valor_base) if valor_base > 0 else 0.0
    deficit_operacional = max(0.0, doa_antes_teto - limite_legal) if valor_base > 0 else 0.0

    return {
        "anos_vigencia": anos,
        "fator_temporal": fator_temporal,
        "fator_governanca": fator_governanca,
        "doa_operacional": total_base,
        "doa_metodologica": doa_metodologica,
        "piso_institucional": piso_institucional,
        "limite_legal": limite_legal,
        "doa_antes_teto": doa_antes_teto,
        "doa_final": doa_final,
        "fator_ajuste_final": fator_ajuste_final,
        "percentual_final": percentual_final,
        "deficit_operacional": deficit_operacional,
    }


def fator_por_nota(nota):
    """Escala 1 a 5: 1=0,90; 2=1,00; 3=1,10; 4=1,20; 5=1,30."""
    try:
        nota = float(nota)
    except Exception:
        nota = 3.0
    nota = max(1.0, min(5.0, nota))
    return 0.8 + (nota * 0.1)


def obter_fatores_complexidade():
    """Calcula fatores de ICP e Giro a partir das notas salvas para a proposta atual."""
    df = carregar_icp_giro()

    if df.empty or not st.session_state.get("id_projeto"):
        return 1.10, 1.10, 1.10, 3.0, 3.0

    filtro = (
        (df["id_projeto"].astype(str) == str(st.session_state.get("id_projeto"))) &
        (df["versao"].astype(int) == int(st.session_state.get("versao")))
    )
    df_atual = df[filtro].copy()

    if df_atual.empty:
        return 1.10, 1.10, 1.10, 3.0, 3.0

    df_atual["nota"] = pd.to_numeric(df_atual["nota"], errors="coerce")

    media_icp = df_atual.loc[df_atual["tipo"].astype(str).str.upper() == "ICP", "nota"].mean()
    media_giro = df_atual.loc[df_atual["tipo"].astype(str).str.upper() == "GIRO", "nota"].mean()

    if pd.isna(media_icp):
        media_icp = 3.0
    if pd.isna(media_giro):
        media_giro = 3.0

    fator_icp = fator_por_nota(media_icp)
    fator_giro = fator_por_nota(media_giro)

    # Peso metodológico inicial: ICP 40% e Giro 60%.
    fator_complexidade = (fator_icp * 0.40) + (fator_giro * 0.60)

    return fator_icp, fator_giro, fator_complexidade, float(media_icp), float(media_giro)




def carregar_explicacoes_indicadores():
    bases = carregar_custos_indicadores()
    return {
        "ICP": bases.get("ICP", pd.DataFrame()).copy(),
        "GIRO": bases.get("GIRO", pd.DataFrame()).copy(),
    }


def explicacao_nota_indicador(tipo, indicador, nota):
    bases = carregar_explicacoes_indicadores()
    df = bases.get(str(tipo).upper(), pd.DataFrame()).copy()

    try:
        nota_num = int(nota)
    except Exception:
        nota_num = 3

    if df.empty:
        return f"Nota {nota_num}: consulte a escala explicativa cadastrada."

    df.columns = df.columns.astype(str).str.strip()
    indicador_txt = normalizar_texto(indicador)

    if str(tipo).upper() == "ICP":
        col_item = "Item" if "Item" in df.columns else df.columns[0]
        col_nota = "Nota" if "Nota" in df.columns else df.columns[1]
        col_desc = "Descrição" if "Descrição" in df.columns else df.columns[-1]

        mapa_icp = {
            "financiador": "Financiador",
            "maturidade": "Planejamento",
            "planejamento": "Planejamento",
            "valor": "Valor Global do Projeto",
            "executor": "Executor",
            "início": "Prazo para Início",
            "inicio": "Prazo para Início",
            "coordenador com a facto": "Experiência do Coordenador com a Facto",
            "coordenador": "Experiência do Coordenador",
            "articulação": "RCI",
            "articulacao": "RCI",
            "sensibilidade": "RCI",
            "reputacional": "RCI",
            "normativa": "Planejamento",
            "jurídica": "Planejamento",
            "juridica": "Planejamento",
            "acompanhamento": "Executor",
            "risco": "Executor",
            "técnico": "Experiência do Coordenador",
            "tecnico": "Experiência do Coordenador",
        }
        item_alvo = None
        for chave, item in mapa_icp.items():
            if chave in indicador_txt:
                item_alvo = item
                break

        temp = df.copy()
        temp[col_nota] = pd.to_numeric(temp[col_nota], errors="coerce")
        candidatos = temp[temp[col_nota] == nota_num]
        if item_alvo:
            candidatos_item = candidatos[candidatos[col_item].astype(str).str.strip().str.casefold() == item_alvo.casefold()]
            if not candidatos_item.empty:
                return f"Nota {nota_num}: {candidatos_item.iloc[0][col_desc]}"
        if not candidatos.empty:
            return f"Nota {nota_num}: {candidatos.iloc[0][col_desc]}"

    if str(tipo).upper() == "GIRO":
        col_item = "GIRO" if "GIRO" in df.columns else df.columns[0]
        col_nota = "Nota" if "Nota" in df.columns else df.columns[1]
        col_desc = "Fator" if "Fator" in df.columns else df.columns[-1]
        temp = df.copy()
        temp[col_nota] = pd.to_numeric(temp[col_nota], errors="coerce")
        candidatos = temp[temp[col_nota] == nota_num]
        candidatos_exato = candidatos[candidatos[col_item].astype(str).str.strip().str.casefold() == str(indicador).strip().casefold()]
        if not candidatos_exato.empty:
            return f"Nota {nota_num}: {candidatos_exato.iloc[0][col_desc]}"
        for _, row in candidatos.iterrows():
            item_norm = normalizar_texto(row[col_item])
            if item_norm and (item_norm in indicador_txt or indicador_txt in item_norm):
                return f"Nota {nota_num}: {row[col_desc]}"
        if not candidatos.empty:
            return f"Nota {nota_num}: {candidatos.iloc[0][col_desc]}"

    return f"Nota {nota_num}: consulte a escala explicativa cadastrada."

def calcular_custo_hora_setor(bases):
    pessoal = bases.get("custos_pessoal", pd.DataFrame()).copy()
    adm = bases.get("custos_adm", pd.DataFrame()).copy()

    if pessoal.empty:
        return pd.DataFrame(columns=["setor", "custo_hora_pessoal", "custo_hora_adm", "custo_hora_total"])

    pessoal.columns = pessoal.columns.astype(str).str.strip()
    pessoal["ch_mensal"] = pd.to_numeric(pessoal.get("ch_mensal"), errors="coerce").fillna(0)
    pessoal["custo_pessoa_mensal"] = pd.to_numeric(pessoal.get("custo_pessoa_mensal"), errors="coerce").fillna(0)

    setor = pessoal.groupby("setor", as_index=False).agg({
        "ch_mensal": "sum",
        "custo_pessoa_mensal": "sum"
    })

    setor["custo_hora_pessoal"] = setor.apply(
        lambda r: r["custo_pessoa_mensal"] / r["ch_mensal"] if r["ch_mensal"] else 0,
        axis=1
    )

    total_horas = setor["ch_mensal"].sum()

    if not adm.empty and "valor" in adm.columns and total_horas:
        adm["valor"] = pd.to_numeric(adm["valor"], errors="coerce").fillna(0)
        custo_hora_adm = adm["valor"].sum() / total_horas
    else:
        custo_hora_adm = 0.0

    setor["custo_hora_adm"] = custo_hora_adm
    setor["custo_hora_total"] = setor["custo_hora_pessoal"] + setor["custo_hora_adm"]

    return setor[["setor", "custo_hora_pessoal", "custo_hora_adm", "custo_hora_total"]]


def adicionar_linha_doa(
    linhas,
    tempo_df,
    custo_hora_df,
    atividade,
    quantidade,
    fator_icp,
    fator_giro,
    fator_complexidade,
    tipo_recorrencia="pontual",
    meses_aplicaveis=1,
    componente_forcado=None,
):
    """Adiciona uma atividade à memória da DOA sem multiplicar tudo pela vigência.

    Regras:
    - pontual: quantidade × tempo-base;
    - ocorrencia: ocorrências totais × tempo-base;
    - mensal: quantidade × tempo-base × meses aplicáveis.

    O parâmetro ``quantidade`` deve representar a unidade operacional real da
    atividade. A vigência só é aplicada quando ``tipo_recorrencia`` for mensal.
    """
    try:
        quantidade = float(quantidade or 0)
    except Exception:
        quantidade = 0.0

    if quantidade <= 0:
        return

    info = buscar_atividade(tempo_df, atividade)
    if not info:
        return

    try:
        meses_aplicaveis = max(1.0, float(meses_aplicaveis or 1))
    except Exception:
        meses_aplicaveis = 1.0

    tipo_recorrencia = str(tipo_recorrencia or "pontual").strip().lower()
    if tipo_recorrencia not in {"pontual", "ocorrencia", "mensal"}:
        tipo_recorrencia = "pontual"

    setor = info.get("setor")
    tempo_base = float(info.get("tempo_base_h") or 0)

    if tipo_recorrencia == "mensal":
        horas_estimadas = quantidade * tempo_base * meses_aplicaveis
        meses_considerados = meses_aplicaveis
    else:
        horas_estimadas = quantidade * tempo_base
        meses_considerados = 1.0

    custo_setor = custo_hora_df[custo_hora_df["setor"].astype(str) == str(setor)]

    if custo_setor.empty:
        custo_hora_pessoal = 0.0
        custo_hora_adm = 0.0
        custo_hora_total = 0.0
    else:
        custo_hora_pessoal = float(custo_setor["custo_hora_pessoal"].iloc[0] or 0)
        custo_hora_adm = float(custo_setor["custo_hora_adm"].iloc[0] or 0)
        custo_hora_total = float(custo_setor["custo_hora_total"].iloc[0] or 0)

    valor_calculado = horas_estimadas * custo_hora_total * fator_complexidade

    linhas.append({
        "componente": componente_forcado or componente_por_setor(setor),
        "setor": setor,
        "atividade": info.get("atividade"),
        "quantidade_operacional": quantidade,
        "tempo_base_h": tempo_base,
        "vigencia_meses": meses_considerados,
        "horas_estimadas": horas_estimadas,
        "custo_hora_pessoal": custo_hora_pessoal,
        "custo_hora_adm": custo_hora_adm,
        "custo_hora_total": custo_hora_total,
        "fator_icp": fator_icp,
        "fator_giro": fator_giro,
        "fator_complexidade": fator_complexidade,
        "valor_calculado": valor_calculado,
    })


def gerar_memoria_doa(df_parf_atual, bases):
    tempo_df = bases.get("tempo_pessoas", pd.DataFrame()).copy()
    if not tempo_df.empty:
        tempo_df.columns = tempo_df.columns.astype(str).str.strip()

    custo_hora_df = calcular_custo_hora_setor(bases)
    fator_icp, fator_giro, fator_complexidade, media_icp, media_giro = obter_fatores_complexidade()

    projeto_atual = obter_projeto_atual() or {}
    try:
        vigencia_projeto = max(1.0, float(projeto_atual.get("prazo_meses", 1) or 1))
    except Exception:
        vigencia_projeto = 1.0

    linhas = []

    for _, row in df_parf_atual.iterrows():
        grupo = str(row.get("grupo", "")).strip()
        modalidade_contratacao = str(row.get("modalidade_contratacao", "")).strip()

        qtd_num = pd.to_numeric(row.get("quantidade", 0), errors="coerce")
        meses_num = pd.to_numeric(row.get("meses", 0), errors="coerce")
        qtd = 0.0 if pd.isna(qtd_num) else float(qtd_num)
        meses = 1.0 if pd.isna(meses_num) or float(meses_num) <= 0 else float(meses_num)

        if grupo == "Bolsas":
            # Uma vez por bolsista.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Implementação de bolsas", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Gestão de contrato de bolsas", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            # Uma ocorrência por bolsista em cada mês de bolsa; meses já entram na quantidade total.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Pagamento de bolsa", qtd * meses, fator_icp, fator_giro, fator_complexidade, "ocorrencia")
            # Uma rotina por mês, independentemente da quantidade de bolsistas.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Gestão de pagamentos de bolsas", 1, fator_icp, fator_giro, fator_complexidade, "mensal", meses)

        elif grupo == "Celetistas":
            # Seleção e contratação são eventos únicos por pessoa.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Seleção de CLT", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Contratação CLT", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            # Gestão mensal por celetista durante os meses informados.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Gestão mensal de contratos (por pessoa)", qtd, fator_icp, fator_giro, fator_complexidade, "mensal", meses)

        elif grupo == "RPA":
            # Cada RPA representa uma ocorrência completa; não há nova multiplicação pela vigência.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Seleção de RPA (Projetos)", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Gestão de Contratos RPA (Projetos)", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Pagamento de RPA", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")

        elif grupo in ["Material de consumo", "Material permanente", "Serviço PJ"]:
            if "Inexigibilidade" in modalidade_contratacao:
                atividade_compra = "Inexigibilidade"
            elif "Licitação" in modalidade_contratacao or "Dispensa" in modalidade_contratacao:
                atividade_compra = "Licitação/Dispensa"
            else:
                atividade_compra = "Compra direta"

            # Processo de contratação e pagamento são contados pela quantidade informada.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, atividade_compra, qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Pagamento de fornecedor", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")

            if grupo == "Material permanente":
                adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Gestão de patrimônio", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")

            if grupo == "Serviço PJ":
                # Gestão contratual mensal durante a vigência do projeto.
                adicionar_linha_doa(
                    linhas, tempo_df, custo_hora_df,
                    "Gestão de contratos de compras e contratações",
                    qtd, fator_icp, fator_giro, fator_complexidade,
                    "mensal", vigencia_projeto
                )

        elif grupo == "Importação":
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Importação", qtd, fator_icp, fator_giro, fator_complexidade, "pontual")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Pagamento de fornecedor", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")

        elif grupo == "Passagens":
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Número de passagens", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Pagamento de fornecedor", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")

        elif grupo == "Diárias":
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Diárias (Projetos)", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Pagamento de diária", qtd, fator_icp, fator_giro, fator_complexidade, "ocorrencia")

        elif grupo == "Tarifas bancárias":
            # Conciliação é uma rotina mensal do projeto.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Conciliação bancária", 1, fator_icp, fator_giro, fator_complexidade, "mensal", vigencia_projeto)

        elif grupo in ["Ressarcimento", "Contrapartida"]:
            # A entrada da receita é pontual; o gerenciamento financeiro é mensal.
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Gerenciamento orçamentário e financeiro (Projetos)", 1, fator_icp, fator_giro, fator_complexidade, "mensal", vigencia_projeto)
            adicionar_linha_doa(linhas, tempo_df, custo_hora_df, "Entrada de receita (notas)", 1, fator_icp, fator_giro, fator_complexidade, "pontual")

        elif grupo == "Prospecção":
            # Uma atividade pontual de articulação institucional.
            adicionar_linha_doa(
                linhas, tempo_df, custo_hora_df, "Comunicação institucional", 1,
                fator_icp, fator_giro, fator_complexidade,
                "pontual", 1,
                componente_forcado="Prospecção e articulação institucional"
            )

    memoria = pd.DataFrame(linhas)

    # Consolida atividades idênticas para evitar linhas repetidas na memória de cálculo.
    if not memoria.empty:
        chaves = [
            "componente", "setor", "atividade", "tempo_base_h", "vigencia_meses",
            "custo_hora_pessoal", "custo_hora_adm", "custo_hora_total",
            "fator_icp", "fator_giro", "fator_complexidade"
        ]
        memoria = memoria.groupby(chaves, as_index=False, dropna=False).agg({
            "quantidade_operacional": "sum",
            "horas_estimadas": "sum",
            "valor_calculado": "sum",
        })

        ordem = [
            "componente", "setor", "atividade", "quantidade_operacional",
            "tempo_base_h", "vigencia_meses", "horas_estimadas",
            "custo_hora_pessoal", "custo_hora_adm", "custo_hora_total",
            "fator_icp", "fator_giro", "fator_complexidade", "valor_calculado"
        ]
        memoria = memoria[ordem]

    return memoria, custo_hora_df, media_icp, media_giro, fator_icp, fator_giro, fator_complexidade


def formatar_tabela_moeda(df, colunas_moeda):
    visual = df.copy()
    for coluna in colunas_moeda:
        if coluna in visual.columns:
            visual[coluna] = visual[coluna].apply(moeda)
    return visual




# =========================
# FUNÇÕES PROPOSTA COMERCIAL
# =========================

def nome_seguro_arquivo(texto):
    texto = str(texto or "proposta").strip().lower()
    substituicoes = {
        "á": "a", "à": "a", "ã": "a", "â": "a",
        "é": "e", "ê": "e",
        "í": "i",
        "ó": "o", "õ": "o", "ô": "o",
        "ú": "u",
        "ç": "c"
    }
    for origem, destino in substituicoes.items():
        texto = texto.replace(origem, destino)
    permitido = []
    for c in texto:
        if c.isalnum():
            permitido.append(c)
        elif c in [" ", "-", "_"]:
            permitido.append("_")
    texto = "".join(permitido)
    while "__" in texto:
        texto = texto.replace("__", "_")
    return texto.strip("_") or "proposta"


def adicionar_paragrafo_apos(paragrafo, texto="", estilo=None):
    from docx.text.paragraph import Paragraph
    from docx.oxml import OxmlElement

    novo_p = OxmlElement("w:p")
    paragrafo._p.addnext(novo_p)
    novo_paragrafo = Paragraph(novo_p, paragrafo._parent)
    if estilo:
        novo_paragrafo.style = estilo
    if texto:
        novo_paragrafo.add_run(texto)
    return novo_paragrafo


def inserir_tabela_apos_paragrafo(documento, paragrafo, dados, estilo="Table Grid"):
    """Insere tabela depois de um parágrafo do modelo e devolve a tabela."""
    linhas = len(dados)
    colunas = len(dados[0]) if dados else 1
    tabela = documento.add_table(rows=linhas, cols=colunas)
    tabela.style = estilo

    for i, linha in enumerate(dados):
        for j, valor in enumerate(linha):
            celula = tabela.cell(i, j)
            celula.text = str(valor if valor is not None else "")
            for par in celula.paragraphs:
                for run in par.runs:
                    run.font.size = None
                    if i == 0:
                        run.bold = True

    paragrafo._p.addnext(tabela._tbl)
    return tabela


def inserir_paragrafo_apos_paragrafo(documento, paragrafo, texto):
    """Insere um parágrafo simples depois de outro parágrafo."""
    novo = documento.add_paragraph(str(texto))
    paragrafo._p.addnext(novo._p)
    return novo


def inserir_paragrafo_antes_paragrafo(documento, paragrafo, texto):
    """Insere um parágrafo simples antes de outro parágrafo."""
    novo = documento.add_paragraph(str(texto))
    paragrafo._p.addprevious(novo._p)
    return novo





def inserir_tabela_antes_paragrafo(documento, paragrafo, dados, estilo="Table Grid"):
    """Insere tabela antes de um parágrafo do modelo e devolve a tabela."""
    linhas = len(dados)
    colunas = len(dados[0]) if dados else 1
    tabela = documento.add_table(rows=linhas, cols=colunas)
    tabela.style = estilo

    for i, linha in enumerate(dados):
        for j, valor in enumerate(linha):
            celula = tabela.cell(i, j)
            celula.text = str(valor if valor is not None else "")
            for par in celula.paragraphs:
                for run in par.runs:
                    if i == 0:
                        run.bold = True

    paragrafo._p.addprevious(tabela._tbl)
    return tabela


def localizar_proximo_paragrafo_nao_vazio(documento, paragrafo_referencia):
    try:
        idx = documento.paragraphs.index(paragrafo_referencia)
    except ValueError:
        return paragrafo_referencia

    for p in documento.paragraphs[idx + 1:]:
        if p.text.strip():
            return p
    return paragrafo_referencia

def localizar_paragrafo_por_texto(documento, texto_busca):
    alvo = str(texto_busca).strip().lower()
    for paragrafo in documento.paragraphs:
        texto = paragrafo.text.strip().lower()
        if alvo in texto:
            return paragrafo
    return None

def inserir_quadro_no_marcador(documento, marcador, titulo, dados):
    """Substitui um marcador textual por legenda + tabela no local exato."""
    par_marcador = localizar_paragrafo_por_texto(documento, marcador)
    if not par_marcador:
        return False

    par_marcador.text = str(titulo)
    inserir_tabela_apos_paragrafo(documento, par_marcador, dados)
    return True



def obter_parf_atual_df():
    df_parf = carregar_parf()
    if df_parf.empty or not id_atual:
        return pd.DataFrame(columns=colunas_parf())
    return df_parf[
        (df_parf["id_projeto"].astype(str) == str(id_atual)) &
        (df_parf["versao"].astype(int) == int(versao_atual))
    ].copy()


def obter_resumo_parf_proposta(df_parf_atual):
    if df_parf_atual.empty:
        return pd.DataFrame(columns=["Rubrica", "Quantidade", "Valor previsto"])

    temp = df_parf_atual.copy()
    temp["quantidade"] = pd.to_numeric(temp["quantidade"], errors="coerce").fillna(0)
    temp["total"] = pd.to_numeric(temp["total"], errors="coerce").fillna(0)

    resumo = temp.groupby("grupo", as_index=False).agg({
        "quantidade": "sum",
        "total": "sum"
    })
    resumo = resumo.rename(columns={
        "grupo": "Rubrica",
        "quantidade": "Quantidade",
        "total": "Valor previsto"
    })
    return resumo


def obter_resumo_doa_proposta(df_parf_atual):
    """Retorna tabela de DOA agrupada por componente. Usa resultado salvo; se não houver, recalcula."""
    df_doa = carregar_doa()
    if not df_doa.empty:
        filtro = (
            (df_doa["id_projeto"].astype(str) == str(id_atual)) &
            (df_doa["versao"].astype(int) == int(versao_atual))
        )
        df_salvo = df_doa[filtro].copy()
        if not df_salvo.empty:
            df_salvo["valor_calculado"] = pd.to_numeric(df_salvo["valor_calculado"], errors="coerce").fillna(0)
            df_salvo["valor_ajustado"] = pd.to_numeric(df_salvo["valor_ajustado"], errors="coerce").fillna(0)
            resumo = df_salvo.groupby("componente", as_index=False).agg({
                "valor_calculado": "sum",
                "valor_ajustado": "sum"
            })
            resumo = resumo.rename(columns={
                "componente": "Componente da DOA",
                "valor_calculado": "Valor calculado",
                "valor_ajustado": "Valor final"
            })
            return resumo

    # Recalcula quando o usuário ainda não clicou em "Salvar resultado da DOA".
    if df_parf_atual.empty or not ARQ_CUSTOS_INDICADORES.exists():
        return pd.DataFrame(columns=["Componente da DOA", "Valor calculado", "Valor final"])

    bases_custos = carregar_custos_indicadores()
    memoria, *_ = gerar_memoria_doa(df_parf_atual, bases_custos)
    if memoria.empty:
        return pd.DataFrame(columns=["Componente da DOA", "Valor calculado", "Valor final"])

    projeto = obter_projeto_atual() or {}
    valor_projeto = float(projeto.get("valor_aprovado", 0) or 0)
    if valor_projeto <= 0:
        valor_projeto = float(pd.to_numeric(df_parf_atual["total"], errors="coerce").fillna(0).sum())

    resumo = memoria.groupby("componente", as_index=False)["valor_calculado"].sum()
    total_calculado = float(resumo["valor_calculado"].sum())
    prazo_meses = float(projeto.get("prazo_meses", 1) or 1)
    ajustes = calcular_ajustes_metodologicos_doa(total_calculado, valor_projeto, prazo_meses)
    resumo["valor_ajustado"] = resumo["valor_calculado"] * ajustes["fator_ajuste_final"]

    resumo = resumo.rename(columns={
        "componente": "Componente da DOA",
        "valor_calculado": "Valor calculado",
        "valor_ajustado": "Valor final"
    })
    return resumo


def gerar_proposta_docx():
    try:
        from docx import Document
    except Exception:
        st.error("Biblioteca python-docx não instalada. Rode: pip install python-docx")
        return None

    if not ARQ_MODELO_PROPOSTA.exists():
        st.error("Modelo não encontrado. Salve o arquivo como dados/modelo_proposta.docx.")
        return None

    projeto = obter_projeto_atual() or {}
    df_parf_atual = obter_parf_atual_df()
    resumo_parf = obter_resumo_parf_proposta(df_parf_atual)
    resumo_doa = obter_resumo_doa_proposta(df_parf_atual)

    documento = Document(str(ARQ_MODELO_PROPOSTA))

    dados_projeto = [
        ["Campo", "Informação"],
        ["Projeto", projeto.get("nome_projeto", "")],
        ["Coordenador", projeto.get("coordenador", "")],
        ["E-mail do coordenador", projeto.get("email_coordenador", "")],
        ["Instituição executora", projeto.get("instituicao_executora", "")],
        ["Campus", projeto.get("campus", "")],
        ["Financiador", projeto.get("financiador", "")],
        ["Tipo de financiador", projeto.get("tipo_financiador", "")],
        ["Tipo de instrumento", projeto.get("tipo_instrumento", "")],
        ["Vigência", f"{projeto.get('data_inicio', '')} a {projeto.get('data_fim', '')}"],
        ["Prazo estimado", f"{projeto.get('prazo_meses', '')} meses"],
        ["Valor aprovado/previsto", moeda(projeto.get("valor_aprovado", 0))],
        ["Área temática", projeto.get("area_tematica", "")],
        ["Responsável interno Facto", projeto.get("responsavel_facto", "")],
    ]

    dados_rubricas = [["Rubrica", "Quantidade", "Valor previsto"]]
    if not resumo_parf.empty:
        for _, row in resumo_parf.iterrows():
            dados_rubricas.append([
                row.get("Rubrica", ""),
                f"{float(row.get('Quantidade', 0) or 0):,.0f}".replace(",", "."),
                moeda(row.get("Valor previsto", 0))
            ])
        dados_rubricas.append(["Total", "", moeda(resumo_parf["Valor previsto"].sum())])
    else:
        dados_rubricas.append(["Sem rubricas cadastradas", "", moeda(0)])

    dados_doa = [["Componente da DOA", "Valor calculado", "Valor final"]]
    if not resumo_doa.empty:
        for _, row in resumo_doa.iterrows():
            dados_doa.append([
                row.get("Componente da DOA", ""),
                moeda(row.get("Valor calculado", 0)),
                moeda(row.get("Valor final", 0))
            ])
        dados_doa.append([
            "Total",
            moeda(resumo_doa["Valor calculado"].sum()),
            moeda(resumo_doa["Valor final"].sum())
        ])
    else:
        dados_doa.append(["DOA ainda não calculada", moeda(0), moeda(0)])

    # Insere quadros exatamente onde os marcadores foram colocados no modelo.
    # Marcadores esperados no modelo_proposta.docx:
    # {{TABELA_DADOS_PROJETO}}
    # {{TABELA_RESUMO_PARF}}
    # {{TABELA_SERVICOS_ITENS}}
    # {{TABELA_DOA}}

    inseriu_dados = inserir_quadro_no_marcador(
        documento,
        "{{TABELA_DADOS_PROJETO}}",
        "Quadro 1 – Dados gerais do projeto",
        dados_projeto
    )

    inseriu_resumo = inserir_quadro_no_marcador(
        documento,
        "{{TABELA_RESUMO_PARF}}",
        "Quadro 2 – Resumo das rubricas previstas no PARF",
        dados_rubricas
    )

    inseriu_servicos = inserir_quadro_no_marcador(
        documento,
        "{{TABELA_SERVICOS_ITENS}}",
        "Quadro 3 – Serviços e itens previstos para execução",
        dados_rubricas
    )

    inseriu_doa = inserir_quadro_no_marcador(
        documento,
        "{{TABELA_DOA}}",
        "Quadro 4 – Detalhamento dos valores finais da DOA",
        dados_doa
    )

    marcadores_faltantes = []
    if not inseriu_dados:
        marcadores_faltantes.append("{{TABELA_DADOS_PROJETO}}")
    if not inseriu_resumo:
        marcadores_faltantes.append("{{TABELA_RESUMO_PARF}}")
    if not inseriu_servicos:
        marcadores_faltantes.append("{{TABELA_SERVICOS_ITENS}}")
    if not inseriu_doa:
        marcadores_faltantes.append("{{TABELA_DOA}}")

    if marcadores_faltantes:
        st.warning(
            "Alguns marcadores não foram encontrados no modelo_proposta.docx: "
            + ", ".join(marcadores_faltantes)
        )

    nome_base = nome_seguro_arquivo(projeto.get("nome_projeto", "proposta"))
    saida = SAIDAS_DIR / f"proposta_{nome_base}_v{versao_atual}_{id_atual}.docx"
    documento.save(str(saida))
    return saida



# =========================
# AUTENTICAÇÃO SIMPLES
# =========================

USUARIOS_AUTORIZADOS = {
    "laiz": "laiz@facto",
    "katarina": "katarina@facto",
    "marcos": "marcos@facto",
    "eliane": "eliane@facto",
    "daiane": "daiane@facto",
    "everton": "everton@facto",
}

if "autenticado" not in st.session_state:
    st.session_state["autenticado"] = False

if "usuario_logado" not in st.session_state:
    st.session_state["usuario_logado"] = ""

if not st.session_state["autenticado"]:
    st.markdown("## Acesso à plataforma")
    st.caption("Informe login e senha para acessar a plataforma de precificação.")

    with st.form("form_login"):
        login = st.text_input("Login").strip().lower()
        senha = st.text_input("Senha", type="password")
        entrar = st.form_submit_button("Entrar")

    if entrar:
        if login in USUARIOS_AUTORIZADOS and senha == USUARIOS_AUTORIZADOS[login]:
            st.session_state["autenticado"] = True
            st.session_state["usuario_logado"] = login
            st.rerun()
        else:
            st.error("Login ou senha inválidos.")

    st.stop()

# =========================
# CABEÇALHO
# =========================

col_logo, col_titulo = st.columns([1, 5])

with col_logo:
    if LOGO_FACTO.exists():
        st.image(str(LOGO_FACTO), width=150)
    else:
        st.markdown("### Facto")

with col_titulo:
    st.markdown(
        f"""
        <div style="
            background: linear-gradient(135deg, {AZUL} 0%, #244a63 100%);
            padding: 24px 32px;
            border-radius: 0 0 24px 24px;
            margin-bottom: 28px;
        ">
            <div class="facto-badge">FACTO • Plataforma Institucional</div>
            <div class="facto-title">Precificação de Projetos e DOA</div>
            <div class="facto-subtitle">
                Cadastro → PARF → ICP/Giro → DOA → Proposta Comercial
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

# =========================
# SIDEBAR
# =========================

st.sidebar.title("Navegação")

st.sidebar.caption(f"Usuário logado: {st.session_state.get("usuario_logado", "")}")
if st.sidebar.button("Sair", use_container_width=True):
    st.session_state["autenticado"] = False
    st.session_state["usuario_logado"] = ""
    st.rerun()


acao = st.sidebar.radio(
    "Ação",
    ["Nova proposta", "Abrir proposta existente"]
)

df_projetos = carregar_projetos()

if acao == "Nova proposta":
    criar = st.sidebar.button(
        "➕ Criar nova proposta",
        use_container_width=True,
        type="primary"
    )

    if criar:
        st.session_state["id_projeto"] = gerar_id()
        st.session_state["versao"] = 1
        st.session_state["status"] = "Rascunho"
        st.rerun()

else:
    if df_projetos.empty:
        st.sidebar.info("Nenhuma proposta cadastrada.")
    else:
        df_projetos = df_projetos.copy()
        df_projetos["label"] = (
            df_projetos["nome_projeto"].fillna("Sem nome").astype(str) +
            " | ID " + df_projetos["id_projeto"].astype(str) +
            " | v" + df_projetos["versao"].astype(str) +
            " | " + df_projetos["status"].astype(str)
        )

        selecionado = st.sidebar.selectbox(
            "Selecione uma proposta",
            df_projetos["label"].tolist()
        )

        linha = df_projetos[df_projetos["label"] == selecionado].iloc[0]

        abrir = st.sidebar.button(
            "Abrir proposta",
            use_container_width=True
        )

        if abrir:
            st.session_state["id_projeto"] = str(linha["id_projeto"])
            st.session_state["versao"] = int(linha["versao"])
            st.session_state["status"] = str(linha["status"])
            st.sidebar.success("Proposta carregada.")
            st.rerun()

        nova_versao_btn = st.sidebar.button(
            "Criar nova versão para renegociação",
            use_container_width=True
        )

        if nova_versao_btn:
            nova = criar_nova_versao(str(linha["id_projeto"]), int(linha["versao"]))

            if nova:
                st.session_state["id_projeto"] = str(linha["id_projeto"])
                st.session_state["versao"] = nova
                st.session_state["status"] = "Renegociada"
                st.sidebar.success(f"Versão {nova} criada.")
                st.rerun()

        st.sidebar.markdown("---")
        confirmar_exclusao = st.sidebar.checkbox(
            "Confirmo a exclusão desta versão",
            value=False,
            key=f"confirmar_exclusao_{linha['id_projeto']}_{linha['versao']}"
        )
        excluir_btn = st.sidebar.button(
            "🗑️ Excluir esta versão da proposta",
            use_container_width=True,
            disabled=not confirmar_exclusao,
            key=f"excluir_proposta_{linha['id_projeto']}_{linha['versao']}"
        )

        if excluir_btn:
            id_excluir = str(linha["id_projeto"])
            versao_excluir = int(linha["versao"])
            excluir_proposta_versao(id_excluir, versao_excluir)

            if (
                str(st.session_state.get("id_projeto", "")) == id_excluir
                and int(st.session_state.get("versao", 0) or 0) == versao_excluir
            ):
                st.session_state.pop("id_projeto", None)
                st.session_state.pop("versao", None)
                st.session_state.pop("status", None)

            st.sidebar.success("Versão da proposta excluída com sucesso.")
            st.rerun()

pagina = st.sidebar.radio(
    "Etapas",
    [
        "1. Cadastro do Projeto",
        "2. PARF",
        "3. ICP e Giro",
        "4. DOA",
        "5. Proposta Comercial",
        "6. Histórico"
    ]
)

id_atual = st.session_state.get("id_projeto")
versao_atual = st.session_state.get("versao")

if id_atual:
    st.sidebar.markdown("---")
    st.sidebar.write(f"**Projeto atual:** {id_atual}")
    st.sidebar.write(f"**Versão:** {versao_atual}")

mostrar_card_proposta()

# =========================
# PÁGINA 1 — CADASTRO
# =========================

if pagina == "1. Cadastro do Projeto":

    st.subheader("1. Cadastro do Projeto")

    if not id_atual:
        st.warning("Crie uma nova proposta ou abra uma proposta existente na barra lateral.")
        st.stop()

    projeto = obter_projeto_atual() or {}

    with st.form("form_cadastro"):
        col1, col2 = st.columns(2)

        with col1:
            nome_projeto = st.text_input("Nome do projeto", value=projeto.get("nome_projeto", ""))
            coordenador = st.text_input("Coordenador", value=projeto.get("coordenador", ""))
            email_coordenador = st.text_input("E-mail do coordenador", value=projeto.get("email_coordenador", ""))
            instituicao_executora = st.text_input("Instituição executora", value=projeto.get("instituicao_executora", ""))
            campus = st.text_input("Campus", value=projeto.get("campus", ""))
            financiador = st.text_input("Financiador", value=projeto.get("financiador", ""))

            tipo_financiador = st.selectbox(
                "Tipo de financiador",
                ["Público Federal", "Público Estadual", "Público Municipal", "Privado", "Organismo Internacional", "Outro"],
                index=0
            )

            tipo_instrumento = st.selectbox(
                "Tipo de instrumento",
                ["Convênio", "Contrato", "TED", "Termo de Cooperação", "Emenda", "Outro"],
                index=0
            )

        with col2:
            data_inicio = st.date_input(
                "Data de início",
                value=parse_data_streamlit(projeto.get("data_inicio"), date.today())
            )
            data_fim = st.date_input(
                "Data de fim",
                value=parse_data_streamlit(projeto.get("data_fim"), date.today())
            )
            prazo_meses = calcular_prazo_meses(data_inicio, data_fim)

            st.info(f"Prazo estimado: {prazo_meses} meses")

            valor_aprovado = st.number_input(
                "Valor aprovado / previsto do projeto",
                min_value=0.0,
                step=1000.0,
                format="%.2f",
                value=float(projeto.get("valor_aprovado", 0) or 0)
            )

            area_tematica = st.selectbox(
                "Área temática",
                ["Ensino", "Pesquisa", "Extensão", "Inovação", "Desenvolvimento institucional", "Concurso/PS", "Outra"],
                index=0
            )

            responsavel_facto = st.text_input("Responsável interno Facto", value=projeto.get("responsavel_facto", ""))

            status = st.selectbox(
                "Status da proposta",
                ["Rascunho", "Enviada", "Não aprovada", "Renegociada", "Aprovada"],
                index=0
            )

        objeto_resumido = st.text_area("Objeto resumido", value=projeto.get("objeto_resumido", ""))

        st.markdown("### Características operacionais previstas")

        c1, c2, c3, c4, c5 = st.columns(5)

        with c1:
            possui_bolsas = st.checkbox("Bolsas/Auxílios", value=bool(projeto.get("possui_bolsas", False)))
        with c2:
            possui_compras = st.checkbox("Compras/Contratações", value=bool(projeto.get("possui_compras", False)))
        with c3:
            possui_concurso_ps = st.checkbox("Concurso/PS", value=bool(projeto.get("possui_concurso_ps", False)))
        with c4:
            possui_clt_rpa_pj = st.checkbox("CLT/RPA/PJ", value=bool(projeto.get("possui_clt_rpa_pj", False)))
        with c5:
            possui_importacao = st.checkbox("Importação", value=bool(projeto.get("possui_importacao", False)))

        salvar = st.form_submit_button("Salvar cadastro")

    if salvar:
        dados = {
            "id_projeto": id_atual,
            "versao": versao_atual,
            "status": status,
            "nome_projeto": nome_projeto,
            "coordenador": coordenador,
            "email_coordenador": email_coordenador,
            "instituicao_executora": instituicao_executora,
            "campus": campus,
            "financiador": financiador,
            "tipo_financiador": tipo_financiador,
            "tipo_instrumento": tipo_instrumento,
            "data_inicio": data_inicio,
            "data_fim": data_fim,
            "prazo_meses": prazo_meses,
            "valor_aprovado": valor_aprovado,
            "objeto_resumido": objeto_resumido,
            "area_tematica": area_tematica,
            "possui_bolsas": possui_bolsas,
            "possui_compras": possui_compras,
            "possui_concurso_ps": possui_concurso_ps,
            "possui_clt_rpa_pj": possui_clt_rpa_pj,
            "possui_importacao": possui_importacao,
            "responsavel_facto": responsavel_facto,
            "data_criacao": projeto.get("data_criacao", agora()),
            "data_atualizacao": agora()
        }

        salvar_ou_atualizar_projeto(dados)
        st.success("Cadastro salvo com sucesso.")
        st.rerun()

# =========================
# PÁGINA 2 — PARF
# =========================

elif pagina == "2. PARF":

    st.subheader("2. PARF — Estrutura Orçamentária")

    if not id_atual:
        st.warning("Crie ou abra uma proposta antes de preencher o PARF.")
        st.stop()

    aux = carregar_auxiliares()
    df_parf = carregar_parf()

    if "parf_form_seq" not in st.session_state:
        st.session_state["parf_form_seq"] = 0

    form_seq = st.session_state["parf_form_seq"]

    grupo = st.selectbox(
        "Grupo de rubrica",
        ["Selecione a rubrica..."] + grupos_parf(),
        index=0,
        key=f"parf_grupo_{form_seq}"
    )

    if grupo == "Selecione a rubrica...":
        st.info("Selecione uma rubrica para iniciar o preenchimento.")
        st.stop()

    modalidade = ""
    categoria = ""
    item = ""
    descricao = ""
    percentual_adicional = 0.0
    valor_adicional = 0.0
    modalidade_contratacao = ""

    with st.container(border=True):
        col1, col2, col3 = st.columns(3)

        with col1:
            valor_sugerido = 0.0

            if grupo == "Bolsas":
                df_bolsas = obter_opcoes_bolsas(aux)

                if not df_bolsas.empty and {"Modalidade", "Categoria"}.issubset(df_bolsas.columns):
                    df_bolsas = df_bolsas.copy()
                    df_bolsas["Modalidade"] = df_bolsas["Modalidade"].astype(str).str.strip()
                    df_bolsas["Categoria"] = df_bolsas["Categoria"].astype(str).str.strip()
                    df_bolsas["Valor"] = pd.to_numeric(df_bolsas["Valor"], errors="coerce").fillna(0)

                    modalidades = sorted(df_bolsas["Modalidade"].dropna().unique().tolist())
                    modalidade = selectbox_obrigatorio(
                        "Modalidade da bolsa",
                        modalidades,
                        key=f"parf_modalidade_bolsa_{form_seq}"
                    )

                    if modalidade:
                        df_filtrado = df_bolsas[
                            df_bolsas["Modalidade"].str.casefold() == str(modalidade).strip().casefold()
                        ].copy()
                        categorias = sorted(df_filtrado["Categoria"].dropna().unique().tolist())
                        categoria = selectbox_obrigatorio(
                            "Categoria da bolsa",
                            categorias,
                            key=f"parf_categoria_bolsa_{nome_seguro_arquivo(modalidade)}_{form_seq}"
                        )

                        if categoria:
                            # O valor da bolsa não é preenchido automaticamente.
                            # A tabela auxiliar serve apenas para orientar modalidade/categoria.
                            valor_sugerido = 0.0
                            item = st.text_input(
                                "Item",
                                value=categoria,
                                key=f"parf_item_bolsa_{nome_seguro_arquivo(categoria)}_{form_seq}"
                            )
                        else:
                            item = st.text_input("Item", value="", key=f"parf_item_bolsa_vazio_{form_seq}")
                    else:
                        item = st.text_input("Item", value="", key=f"parf_item_bolsa_sem_modalidade_{form_seq}")

                    descricao = st.text_input("Descrição", value="Bolsa", key=f"parf_desc_bolsa_{form_seq}")

                else:
                    modalidade = st.text_input("Modalidade da bolsa", key=f"modalidade_bolsa_manual_{form_seq}")
                    categoria = st.text_input("Categoria da bolsa", key=f"categoria_bolsa_manual_{form_seq}")
                    item = st.text_input("Item", key=f"item_bolsa_manual_{form_seq}")
                    descricao = st.text_input("Descrição", value="Bolsa", key=f"desc_bolsa_manual_{form_seq}")

            elif grupo in grupos_contratacao():
                modalidade = grupo
                st.text_input("Modalidade", value=grupo, disabled=True, key=f"parf_modalidade_{nome_seguro_arquivo(grupo)}_{form_seq}")

                opcoes = obter_lista_despesas(aux, grupo)
                categoria = selectbox_obrigatorio(
                    "Categoria",
                    opcoes,
                    key=f"parf_categoria_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )

                item = st.text_input("Item específico", value="", key=f"parf_item_{nome_seguro_arquivo(grupo)}_{form_seq}")
                descricao = st.text_input("Descrição", value="", key=f"parf_desc_{nome_seguro_arquivo(grupo)}_{form_seq}")

            elif grupo == "Diárias":
                df_diarias = obter_lista_diarias(aux)
                opcoes = df_diarias["tipo"].tolist() if not df_diarias.empty else []

                modalidade = selectbox_obrigatorio(
                    "Tipo de diária/auxílio",
                    opcoes,
                    key=f"parf_tipo_diaria_{form_seq}"
                )

                categoria = "Diárias e auxílio financeiro para viagem"
                item = st.text_input("Item", value=modalidade, key=f"parf_item_diaria_{nome_seguro_arquivo(modalidade)}_{form_seq}")
                descricao = st.text_input("Descrição", value="Diária/Auxílio financeiro para viagem", key=f"parf_desc_diaria_{form_seq}")

                if modalidade and not df_diarias.empty:
                    valores = df_diarias.loc[
                        df_diarias["tipo"].astype(str).str.strip().str.casefold() == str(modalidade).strip().casefold(),
                        "valor"
                    ]
                    if not valores.empty:
                        valor_sugerido = float(valores.iloc[0])

            elif grupo == "Passagens":
                opcoes = obter_lista_passagens(aux)

                modalidade = selectbox_obrigatorio(
                    "Tipo de passagem",
                    opcoes,
                    key=f"parf_tipo_passagem_{form_seq}"
                )

                categoria = "Passagens"
                item = st.text_input("Item", value=modalidade, key=f"parf_item_passagem_{nome_seguro_arquivo(modalidade)}_{form_seq}")
                descricao = st.text_input("Descrição", value="Passagem", key=f"parf_desc_passagem_{form_seq}")

            elif grupo == "Celetistas":
                modalidade = st.selectbox(
                    "Modalidade",
                    ["Selecione...", "CLT - Celetista"],
                    index=0,
                    key=f"parf_modalidade_clt_{form_seq}"
                )
                if modalidade == "Selecione...":
                    modalidade = ""
                categoria = st.text_input("Cargo/Função", value="", key=f"parf_cargo_clt_{form_seq}")
                item = st.text_input("Item", value=categoria, key=f"parf_item_clt_{form_seq}")
                descricao = st.text_input("Descrição", value="Contratação celetista", key=f"parf_desc_clt_{form_seq}")

            elif grupo == "Importação":
                modalidade = st.selectbox(
                    "Modalidade de despesas acessórias",
                    ["Selecione...", "Importação padrão — 60%", "Importação reduzida — 20%"],
                    index=0,
                    key=f"parf_modalidade_importacao_{form_seq}"
                )
                if modalidade == "Importação padrão — 60%":
                    percentual_adicional = 60.0
                elif modalidade == "Importação reduzida — 20%":
                    percentual_adicional = 20.0
                else:
                    modalidade = ""
                    percentual_adicional = 0.0

                categoria = "Importação"
                item = st.text_input("Item importado", value="", key=f"parf_item_importacao_{form_seq}")
                descricao = st.text_input("Descrição", value="Importação", key=f"parf_desc_importacao_{form_seq}")

            elif grupo in ["Ressarcimento", "Contrapartida", "Prospecção"]:
                modalidade = grupo
                categoria = grupo
                item = st.text_input("Item", value=grupo, key=f"parf_item_percentual_{nome_seguro_arquivo(grupo)}_{form_seq}")
                descricao = st.text_input("Descrição", value=f"{grupo} calculado sobre o valor aprovado/previsto", key=f"parf_desc_percentual_{nome_seguro_arquivo(grupo)}_{form_seq}")

            elif grupo == "Tarifas bancárias":
                modalidade = "Tarifas bancárias"
                categoria = "Tarifas bancárias e manutenção de conta"
                item = "Tarifas bancárias"
                descricao = "Rubrica fixa: tarifas por operação bancária e manutenção mensal de conta"
                st.markdown(
                    f"""
                    <div class="facto-info-card">
                        <strong>Rubrica fixa do projeto</strong><br>
                        Calculada automaticamente a partir das operações já lançadas no PARF e da vigência cadastrada.
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            else:
                modalidade = st.text_input("Modalidade", value="", key=f"parf_modalidade_generica_{nome_seguro_arquivo(grupo)}_{form_seq}")
                categoria = st.text_input("Categoria", value="", key=f"parf_categoria_generica_{nome_seguro_arquivo(grupo)}_{form_seq}")
                item = st.text_input("Item", value="", key=f"parf_item_generico_{nome_seguro_arquivo(grupo)}_{form_seq}")
                descricao = st.text_input("Descrição", value="", key=f"parf_desc_generico_{nome_seguro_arquivo(grupo)}_{form_seq}")

        with col2:
            projeto_atual_form = obter_projeto_atual() or {}
            vigencia_projeto_form = float(projeto_atual_form.get("prazo_meses", 0) or 0)

            if grupo in ["Bolsas", "Celetistas", "RPA"]:
                quantidade = st.number_input(
                    "Quantidade",
                    min_value=0.0,
                    step=1.0,
                    value=1.0,
                    key=f"parf_quantidade_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )
                meses = st.number_input(
                    "Meses",
                    min_value=0.0,
                    step=1.0,
                    value=1.0,
                    key=f"parf_meses_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )
                if vigencia_projeto_form and meses > vigencia_projeto_form:
                    st.warning(
                        f"Atenção: o tempo informado ({meses:.0f} meses) é maior que a vigência do projeto "
                        f"({vigencia_projeto_form:.0f} meses). Ajuste a vigência da rubrica ou revise o cadastro do projeto."
                    )

            elif grupo in ["Material de consumo", "Material permanente", "Serviço PJ", "Importação", "Diárias", "Passagens"]:
                quantidade = st.number_input(
                    "Quantidade",
                    min_value=0.0,
                    step=1.0,
                    value=1.0,
                    key=f"parf_quantidade_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )
                meses = 1.0

            elif grupo == "Tarifas bancárias":
                df_parf_base = carregar_parf()
                df_parf_atual_base = df_parf_base[
                    (df_parf_base["id_projeto"].astype(str) == str(id_atual)) &
                    (df_parf_base["versao"].astype(int) == int(versao_atual))
                ].copy() if not df_parf_base.empty else pd.DataFrame()
                vigencia_calc = obter_vigencia_projeto_atual()
                operacoes_tarifa, valor_operacoes_tarifa, valor_manutencao_tarifa, valor_sugerido = calcular_tarifas_bancarias_estimadas(df_parf_atual_base, vigencia_calc)
                quantidade = 1.0
                meses = 1.0
                st.markdown(
                    f"""
                    <div class="facto-info-card">
                        <strong>Composição automática</strong><br>
                        Operações estimadas: <strong>{operacoes_tarifa:.0f}</strong><br>
                        Tarifas por operação: {moeda(valor_operacoes_tarifa)}<br>
                        Manutenção de conta: {moeda(valor_manutencao_tarifa)} ({moeda(100)} × {vigencia_calc:.0f} meses)<br>
                        <strong>Total calculado: {moeda(valor_sugerido)}</strong>
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            elif grupo in ["Ressarcimento", "Contrapartida", "Prospecção"]:
                quantidade = 1.0
                meses = 1.0
                percentual_aplicado = st.number_input(
                    "Percentual sobre o valor aprovado/previsto (%)",
                    min_value=0.0,
                    max_value=100.0,
                    step=0.1,
                    value=0.0,
                    format="%.2f",
                    key=f"parf_percentual_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )
                valor_aprovado_base = float((obter_projeto_atual() or {}).get("valor_aprovado", 0) or 0)
                valor_sugerido = valor_aprovado_base * (percentual_aplicado / 100)
                st.markdown(
                    f"""
                    <div class="facto-info-card">
                        <strong>Valor calculado:</strong> {moeda(valor_sugerido)}<br>
                        Base de cálculo: {moeda(valor_aprovado_base)} × {percentual_aplicado:.2f}%
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            else:
                quantidade = 1.0
                meses = 1.0
                st.info("Para esta rubrica, informe o valor total previsto no campo Valor unitário.")

        with col3:
            if grupo in ["Ressarcimento", "Contrapartida", "Prospecção", "Tarifas bancárias"]:
                valor_unitario = float(valor_sugerido or 0)
                st.markdown(
                    f"""
                    <div class="metric-card">
                        <div class="metric-label">Valor calculado da rubrica</div>
                        <div class="metric-value">{moeda(valor_unitario)}</div>
                    </div>
                    """,
                    unsafe_allow_html=True
                )
            else:
                valor_unitario = st.number_input(
                    "Valor unitário",
                    min_value=0.0,
                    step=100.0,
                    format="%.2f",
                    value=float(valor_sugerido or 0),
                    key=f"parf_valor_unitario_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )

            if grupo == "RPA":
                percentual_adicional = 20.0
                st.info("INSS patronal considerado: 20%")

            elif grupo == "Importação":
                if percentual_adicional:
                    st.info(f"Despesas acessórias de importação consideradas: {percentual_adicional:.0f}%")
                else:
                    st.warning("Selecione a modalidade de despesas acessórias da importação.")

            elif grupo == "Tarifas bancárias":
                st.markdown(
                    """
                    <div class="facto-info-card">
                        <strong>Parâmetros:</strong> R$ 2,06 por operação bancária e R$ 100,00 por mês de manutenção de conta.
                    </div>
                    """,
                    unsafe_allow_html=True
                )

            if grupo in grupos_contratacao():
                marcacao = st.radio(
                    "Tratamento da contratação",
                    ["Automático", "Dispensa", "Inexigibilidade"],
                    horizontal=True,
                    key=f"parf_tratamento_{nome_seguro_arquivo(grupo)}_{form_seq}"
                )
            else:
                marcacao = "Automático"

        subtotal = quantidade * max(1, meses) * valor_unitario
        valor_adicional = subtotal * (percentual_adicional / 100)
        total = subtotal + valor_adicional

        if grupo in grupos_contratacao():
            if marcacao == "Dispensa":
                modalidade_contratacao = "Dispensa"
            elif marcacao == "Inexigibilidade":
                modalidade_contratacao = "Inexigibilidade"
            elif total > 40000:
                modalidade_contratacao = "Licitação"
            else:
                modalidade_contratacao = "Compra direta"
        else:
            modalidade_contratacao = ""

        st.markdown(f"**Total estimado do item:** {moeda(total)}")

        if modalidade_contratacao:
            st.markdown(f"**Modalidade estimada:** {modalidade_contratacao}")

        salvar_item = st.button("Adicionar item ao PARF", type="primary", key=f"salvar_item_parf_{form_seq}")

    if salvar_item:
        if grupo in ["Ressarcimento", "Contrapartida", "Prospecção"] and valor_unitario <= 0:
            st.warning("Informe um percentual maior que zero para incluir esta rubrica no PARF.")
        elif grupo == "Tarifas bancárias" and valor_unitario <= 0:
            st.warning("A tarifa bancária só poderá ser lançada após haver vigência cadastrada ou operações estimadas.")
        elif grupo in ["Bolsas", "Material de consumo", "Material permanente", "Serviço PJ", "Diárias", "Passagens", "Importação", "Celetistas"] and not modalidade:
            st.warning("Selecione/preencha a modalidade antes de adicionar o item.")
        else:
            novo = {
                "id_item": gerar_id(),
                "id_projeto": id_atual,
                "versao": versao_atual,
                "grupo": grupo,
                "item": item,
                "modalidade": modalidade,
                "categoria": categoria,
                "descricao": descricao,
                "quantidade": quantidade,
                "meses": meses,
                "valor_unitario": valor_unitario,
                "percentual_adicional": percentual_adicional,
                "valor_adicional": valor_adicional,
                "modalidade_contratacao": modalidade_contratacao,
                "total": total
            }

            df_parf = pd.concat([df_parf, pd.DataFrame([novo])], ignore_index=True)
            salvar_excel(df_parf, ARQ_PARF)
            st.session_state["parf_form_seq"] += 1
            st.success("Item adicionado ao PARF.")
            st.rerun()

    st.divider()

    df_parf = carregar_parf()

    if df_parf.empty:
        df_atual = pd.DataFrame(columns=colunas_parf())
    else:
        df_atual = df_parf[
            (df_parf["id_projeto"].astype(str) == str(id_atual)) &
            (df_parf["versao"].astype(int) == int(versao_atual))
        ].copy()

    if df_atual.empty:
        st.info("Nenhum item cadastrado no PARF ainda.")
    else:
        st.subheader("Itens do PARF")

        df_atual["total"] = pd.to_numeric(df_atual["total"], errors="coerce").fillna(0)
        resumo = df_atual.groupby("grupo", as_index=False)["total"].sum()
        total_execucao = resumo["total"].sum()

        df_visual = df_atual.rename(columns={
            "grupo": "Grupo",
            "item": "Item",
            "modalidade": "Modalidade",
            "categoria": "Categoria",
            "descricao": "Descrição",
            "quantidade": "Quantidade",
            "meses": "Meses",
            "valor_unitario": "Valor unitário",
            "percentual_adicional": "% adicional",
            "valor_adicional": "Valor adicional",
            "modalidade_contratacao": "Modalidade de contratação",
            "total": "Total"
        })

        colunas_exibir = [
            "Grupo", "Item", "Modalidade", "Categoria", "Descrição",
            "Quantidade", "Meses", "Valor unitário", "% adicional",
            "Valor adicional", "Modalidade de contratação", "Total"
        ]

        df_visual_total = df_visual[colunas_exibir].copy()
        for col_num in ["Quantidade", "Valor unitário", "Valor adicional", "Total"]:
            if col_num in df_visual_total.columns:
                df_visual_total[col_num] = pd.to_numeric(df_visual_total[col_num], errors="coerce").fillna(0)

        linha_total = {col: "" for col in colunas_exibir}
        linha_total["Grupo"] = "Total"
        linha_total["Quantidade"] = df_visual_total["Quantidade"].sum() if "Quantidade" in df_visual_total.columns else 0
        linha_total["Valor unitário"] = df_visual_total["Valor unitário"].sum() if "Valor unitário" in df_visual_total.columns else 0
        linha_total["Valor adicional"] = df_visual_total["Valor adicional"].sum() if "Valor adicional" in df_visual_total.columns else 0
        linha_total["Total"] = df_visual_total["Total"].sum() if "Total" in df_visual_total.columns else 0
        df_visual_total = pd.concat([df_visual_total, pd.DataFrame([linha_total])], ignore_index=True)

        for coluna in ["Valor unitário", "Valor adicional", "Total"]:
            df_visual_total[coluna] = df_visual_total[coluna].apply(moeda)

        st.dataframe(df_visual_total, use_container_width=True, hide_index=True)

        col1, col2 = st.columns([2, 1])

        with col1:
            resumo_visual = resumo.rename(columns={"grupo": "Grupo", "total": "Total"})
            resumo_visual["Total"] = resumo_visual["Total"].apply(moeda)
            st.markdown("### Resumo por rubrica")
            st.dataframe(resumo_visual, use_container_width=True, hide_index=True)

        with col2:
            st.markdown(
                f"""
                <div class="metric-card">
                    <div class="metric-label">Valor total de execução do projeto</div>
                    <div class="metric-value">{moeda(total_execucao)}</div>
                </div>
                """,
                unsafe_allow_html=True
            )

        st.markdown("### Composição do valor do projeto")
        projeto_atual = obter_projeto_atual() or {}
        valor_total_projeto = float(projeto_atual.get("valor_aprovado", 0) or 0)

        try:
            resumo_doa_comp = obter_resumo_doa_proposta(df_atual)
            doa_final = float(pd.to_numeric(resumo_doa_comp.get("Valor final", pd.Series(dtype=float)), errors="coerce").fillna(0).sum())
        except Exception:
            doa_final = 0.0

        valor_com_doa = total_execucao + doa_final
        saldo_disponivel_para_execucao = valor_total_projeto - doa_final if valor_total_projeto else 0.0
        diferenca_total = valor_total_projeto - valor_com_doa if valor_total_projeto else 0.0

        comp1, comp2, comp3, comp4 = st.columns(4)
        comp1.metric("Rubricas de execução", moeda(total_execucao))
        comp2.metric("DOA estimada/final", moeda(doa_final))
        comp3.metric("Execução + DOA", moeda(valor_com_doa))
        comp4.metric("Valor aprovado/previsto", moeda(valor_total_projeto))

        st.markdown(
            f"""
            <div class="facto-callout">
                <strong>Leitura gerencial:</strong>
                o valor aprovado do projeto precisa comportar as rubricas de execução e a DOA.
                Com a DOA estimada, o valor disponível para execução técnica seria <strong>{moeda(saldo_disponivel_para_execucao)}</strong>.
                Diferença entre valor aprovado e composição atual: <strong>{moeda(diferenca_total)}</strong>.
            </div>
            """,
            unsafe_allow_html=True
        )

        st.divider()
        st.subheader("Editar ou excluir item do PARF")

        df_edicao = df_atual.copy()
        if "id_item" not in df_edicao.columns:
            df_edicao["id_item"] = [gerar_id() for _ in range(len(df_edicao))]

        df_edicao["label_edicao"] = (
            df_edicao["grupo"].fillna("").astype(str) + " | " +
            df_edicao["item"].fillna("Sem item").astype(str) + " | " +
            df_edicao["total"].apply(moeda)
        )

        item_escolhido = st.selectbox("Selecione um item salvo", df_edicao["label_edicao"].tolist())
        linha_item = df_edicao[df_edicao["label_edicao"] == item_escolhido].iloc[0]

        with st.form("form_editar_item_parf"):
            col_e1, col_e2, col_e3 = st.columns(3)
            grupos_lista = grupos_parf()
            grupo_atual = str(linha_item.get("grupo", "Bolsas") or "Bolsas")
            index_grupo = grupos_lista.index(grupo_atual) if grupo_atual in grupos_lista else 0

            with col_e1:
                novo_grupo = st.selectbox("Grupo", grupos_lista, index=index_grupo)
                nova_modalidade = st.text_input("Modalidade", value=str(linha_item.get("modalidade", "") or ""))
                nova_categoria = st.text_input("Categoria", value=str(linha_item.get("categoria", "") or ""))
                novo_item = st.text_input("Item", value=str(linha_item.get("item", "") or ""))
                nova_descricao = st.text_input("Descrição", value=str(linha_item.get("descricao", "") or ""))

            with col_e2:
                nova_quantidade = st.number_input("Quantidade", min_value=0.0, step=1.0, value=float(linha_item.get("quantidade", 0) or 0))
                novos_meses = st.number_input("Meses", min_value=0.0, step=1.0, value=float(linha_item.get("meses", 0) or 0))
                novo_valor_unitario = st.number_input(
                    "Valor unitário", min_value=0.0, step=100.0, format="%.2f",
                    value=float(linha_item.get("valor_unitario", 0) or 0)
                )

            with col_e3:
                novo_percentual_adicional = st.number_input("% adicional", min_value=0.0, step=1.0, value=float(linha_item.get("percentual_adicional", 0) or 0))
                nova_modalidade_contratacao = st.text_input("Modalidade de contratação", value=str(linha_item.get("modalidade_contratacao", "") or ""))

            novo_subtotal = nova_quantidade * max(1, novos_meses) * novo_valor_unitario
            novo_valor_adicional = novo_subtotal * (novo_percentual_adicional / 100)
            novo_total = novo_subtotal + novo_valor_adicional

            st.markdown(f"**Novo total estimado:** {moeda(novo_total)}")

            col_btn1, col_btn2 = st.columns(2)
            with col_btn1:
                salvar_edicao = st.form_submit_button("Salvar alteração")
            with col_btn2:
                excluir_item = st.form_submit_button("Excluir item")

        if salvar_edicao:
            dados_atualizados = {
                "grupo": novo_grupo,
                "item": novo_item,
                "modalidade": nova_modalidade,
                "categoria": nova_categoria,
                "descricao": nova_descricao,
                "quantidade": nova_quantidade,
                "meses": novos_meses,
                "valor_unitario": novo_valor_unitario,
                "percentual_adicional": novo_percentual_adicional,
                "valor_adicional": novo_valor_adicional,
                "modalidade_contratacao": nova_modalidade_contratacao,
                "total": novo_total
            }
            ok = atualizar_item_parf(linha_item["id_item"], dados_atualizados)
            if ok:
                st.success("Item atualizado com sucesso.")
                st.rerun()
            else:
                st.error("Não foi possível atualizar o item.")

        if excluir_item:
            ok = excluir_item_parf(linha_item["id_item"])
            if ok:
                st.success("Item excluído com sucesso.")
                st.rerun()
            else:
                st.error("Não foi possível excluir o item.")

# =========================
# PÁGINA 3 — ICP/GIRO
# =========================

elif pagina == "3. ICP e Giro":

    st.subheader("3. ICP e Giro")

    if not id_atual:
        st.warning("Crie ou abra uma proposta antes de preencher ICP/Giro.")
        st.stop()

    escala = {
        1: "Muito baixa complexidade",
        2: "Baixa complexidade",
        3: "Média complexidade",
        4: "Alta complexidade",
        5: "Muito alta complexidade"
    }

    indicadores_icp = [
        "Complexidade do financiador",
        "Maturidade do projeto",
        "Dependência de articulação institucional",
        "Sensibilidade institucional/reputacional",
        "Complexidade normativa/jurídica",
        "Intensidade de acompanhamento gerencial",
        "Risco operacional",
        "Dependência de conhecimento técnico especializado"
    ]

    indicadores_giro = [
        "Giro em compras e contratações",
        "Giro em bolsas, auxílios e seleções",
        "Giro em projetos",
        "Giro administrativo e financeiro",
        "Giro em gestão institucional e governança"
    ]

    with st.expander("Ver escala explicativa geral"):
        for nota, texto in escala.items():
            st.write(f"**{nota}** — {texto}")
        st.caption("Além da escala geral, passe o mouse sobre o ícone de ajuda de cada indicador para ver a explicação da nota selecionada.")

    registros = []

    st.markdown("### ICP")

    for indicador in indicadores_icp:
        col1, col2 = st.columns([1, 2])
        with col1:
            chave = f"icp_{indicador}"
            valor_atual = st.session_state.get(chave, 3)
            help_texto = explicacao_nota_indicador("ICP", indicador, valor_atual)
            nota = st.slider(indicador, 1, 5, int(valor_atual), key=chave, help=help_texto)
            st.caption(explicacao_nota_indicador("ICP", indicador, nota))
        with col2:
            justificativa = st.text_input(f"Justificativa — {indicador}", key=f"just_icp_{indicador}")

        registros.append({
            "id_projeto": id_atual,
            "versao": versao_atual,
            "tipo": "ICP",
            "indicador": indicador,
            "nota": nota,
            "justificativa": justificativa
        })

    st.markdown("### Giro")

    for indicador in indicadores_giro:
        col1, col2 = st.columns([1, 2])
        with col1:
            chave = f"giro_{indicador}"
            valor_atual = st.session_state.get(chave, 3)
            help_texto = explicacao_nota_indicador("GIRO", indicador, valor_atual)
            nota = st.slider(indicador, 1, 5, int(valor_atual), key=chave, help=help_texto)
            st.caption(explicacao_nota_indicador("GIRO", indicador, nota))
        with col2:
            justificativa = st.text_input(f"Justificativa — {indicador}", key=f"just_giro_{indicador}")

        registros.append({
            "id_projeto": id_atual,
            "versao": versao_atual,
            "tipo": "GIRO",
            "indicador": indicador,
            "nota": nota,
            "justificativa": justificativa
        })

    if st.button("Salvar ICP e Giro"):
        df = carregar_icp_giro()

        if not df.empty:
            df = df[
                ~(
                    (df["id_projeto"].astype(str) == str(id_atual)) &
                    (df["versao"].astype(int) == int(versao_atual))
                )
            ]

        df = pd.concat([df, pd.DataFrame(registros)], ignore_index=True)
        salvar_excel(df, ARQ_ICP_GIRO)
        st.success("ICP e Giro salvos com sucesso.")

    df_medias = pd.DataFrame(registros)

    col_a, col_b = st.columns(2)
    with col_a:
        st.metric("ICP médio", round(df_medias[df_medias["tipo"] == "ICP"]["nota"].mean(), 2))
    with col_b:
        st.metric("Giro médio", round(df_medias[df_medias["tipo"] == "GIRO"]["nota"].mean(), 2))

# =========================
# PÁGINA 4 — DOA
# =========================

elif pagina == "4. DOA":

    st.subheader("4. DOA — Cálculo Operacional")

    if not id_atual:
        st.warning("Crie ou abra uma proposta antes de calcular a DOA.")
        st.stop()

    if not ARQ_CUSTOS_INDICADORES.exists():
        st.error("Arquivo de custos não encontrado. Salve a planilha como dados/custos_indicadores.xlsx.")
        st.stop()

    bases_custos = carregar_custos_indicadores()

    df_parf = carregar_parf()
    df_parf_atual = df_parf[
        (df_parf["id_projeto"].astype(str) == str(id_atual)) &
        (df_parf["versao"].astype(int) == int(versao_atual))
    ].copy()

    if df_parf_atual.empty:
        st.warning("Nenhum item foi lançado no PARF desta proposta. Preencha o PARF antes de calcular a DOA.")
        st.stop()

    projeto = obter_projeto_atual() or {}
    valor_projeto = float(projeto.get("valor_aprovado", 0) or 0)

    if valor_projeto <= 0:
        valor_projeto = float(pd.to_numeric(df_parf_atual["total"], errors="coerce").fillna(0).sum())
        st.info("Como o valor aprovado/previsto do projeto não foi informado, o limite de 15% foi calculado sobre o total do PARF.")

    memoria, custo_hora_df, media_icp, media_giro, fator_icp, fator_giro, fator_complexidade = gerar_memoria_doa(df_parf_atual, bases_custos)

    if memoria.empty:
        st.error("Não foi possível gerar memória de cálculo. Verifique se a aba tempo_pessoas contém as atividades esperadas.")
        st.stop()

    resumo_componentes = memoria.groupby("componente", as_index=False)["valor_calculado"].sum()

    total_calculado = float(resumo_componentes["valor_calculado"].sum())
    total_parf_execucao = float(pd.to_numeric(df_parf_atual["total"], errors="coerce").fillna(0).sum())
    prazo_meses = float(projeto.get("prazo_meses", 1) or 1)
    ajustes_doa = calcular_ajustes_metodologicos_doa(total_calculado, valor_projeto, prazo_meses)

    fator_limitador = ajustes_doa["fator_ajuste_final"]
    total_ajustado = ajustes_doa["doa_final"]
    limite_legal = ajustes_doa["limite_legal"]
    percentual_calculado = ajustes_doa["percentual_final"]
    deficit_operacional = ajustes_doa["deficit_operacional"]

    resumo_componentes["valor_ajustado"] = resumo_componentes["valor_calculado"] * fator_limitador

    st.info(
        "Metodologia: as atividades pontuais e por ocorrência não são multiplicadas pela vigência; "
        "os meses são aplicados apenas às rotinas mensais. Sobre a DOA operacional são aplicados, "
        "automaticamente, um fator amortecido de sustentação temporal e um fator de governança "
        "associado ao porte financeiro do projeto. A DOA final respeita piso institucional de 5% "
        "e teto legal de 15% do valor aprovado/previsto."
    )

    st.markdown("### Resultado da DOA")

    card_cols = st.columns(5)
    cards = [
        ("DOA operacional", moeda(total_calculado)),
        ("DOA metodológica", moeda(ajustes_doa["doa_metodologica"])),
        ("Piso institucional (5%)", moeda(ajustes_doa["piso_institucional"])),
        ("DOA final", moeda(total_ajustado)),
        ("Percentual final", f"{percentual_calculado:.2%}".replace(".", ",")),
    ]
    for col, (label, valor) in zip(card_cols, cards):
        with col:
            st.markdown(
                f"""
                <div class="metric-card">
                    <div class="metric-label">{label}</div>
                    <div class="metric-value">{valor}</div>
                </div>
                """,
                unsafe_allow_html=True
            )

    with st.expander("Ver parâmetros automáticos da metodologia"):
        col_m1, col_m2, col_m3, col_m4 = st.columns(4)
        with col_m1:
            st.metric("Vigência equivalente", f"{ajustes_doa['anos_vigencia']:.2f} anos".replace(".", ","))
        with col_m2:
            st.metric("Fator temporal", f"{ajustes_doa['fator_temporal']:.2f}".replace(".", ","))
        with col_m3:
            st.metric("Fator de governança", f"{ajustes_doa['fator_governanca']:.2f}".replace(".", ","))
        with col_m4:
            st.metric("Limite legal (15%)", moeda(ajustes_doa["limite_legal"]))
        if deficit_operacional > 0:
            st.warning(
                "A necessidade metodológica ultrapassou o teto legal. "
                f"Déficit operacional não coberto: {moeda(deficit_operacional)}."
            )

    st.markdown("### Fatores ICP e Giro aplicados")

    col_i1, col_i2, col_i3 = st.columns(3)
    with col_i1:
        st.metric("ICP médio", round(media_icp, 2))
        st.caption(f"Fator ICP aplicado: {fator_icp:.2f}".replace(".", ","))
    with col_i2:
        st.metric("Giro médio", round(media_giro, 2))
        st.caption(f"Fator Giro aplicado: {fator_giro:.2f}".replace(".", ","))
    with col_i3:
        st.metric("Fator ponderado", f"{fator_complexidade:.2f}".replace(".", ","))
        st.caption("Peso inicial: ICP 40% e Giro 60%")

    st.markdown("### Composição da DOA por componente")

    resumo_tabela = resumo_componentes.copy()
    total_row = {
        "componente": "Total",
        "valor_calculado": resumo_tabela["valor_calculado"].sum(),
        "valor_ajustado": resumo_tabela["valor_ajustado"].sum(),
    }
    resumo_tabela = pd.concat([resumo_tabela, pd.DataFrame([total_row])], ignore_index=True)

    resumo_visual = resumo_tabela.rename(columns={
        "componente": "Componente da DOA",
        "valor_calculado": "Valor calculado",
        "valor_ajustado": "Valor final da DOA"
    })
    resumo_visual = formatar_tabela_moeda(resumo_visual, ["Valor calculado", "Valor final da DOA"])

    st.dataframe(resumo_visual, use_container_width=True, hide_index=True)

    with st.expander("Ver memória de cálculo detalhada"):
        memoria_visual = memoria.rename(columns={
            "componente": "Componente",
            "setor": "Setor",
            "atividade": "Atividade",
            "quantidade_operacional": "Quantidade operacional",
            "tempo_base_h": "Tempo base (h)",
            "vigencia_meses": "Meses aplicáveis",
            "horas_estimadas": "Horas estimadas",
            "custo_hora_pessoal": "Custo hora pessoal",
            "custo_hora_adm": "Custo hora administrativo",
            "custo_hora_total": "Custo hora total",
            "fator_complexidade": "Fator ICP/Giro",
            "valor_calculado": "Valor calculado"
        })

        memoria_visual = formatar_tabela_moeda(
            memoria_visual,
            ["Custo hora pessoal", "Custo hora administrativo", "Custo hora total", "Valor calculado"]
        )

        st.dataframe(memoria_visual, use_container_width=True, hide_index=True)

    with st.expander("Explicações das faixas ICP e Giro escolhidas"):
        df_indicadores = carregar_icp_giro()
        if not df_indicadores.empty:
            filtro_ind = (
                (df_indicadores["id_projeto"].astype(str) == str(id_atual)) &
                (df_indicadores["versao"].astype(int) == int(versao_atual))
            )
            df_ind_atual = df_indicadores[filtro_ind].copy()
        else:
            df_ind_atual = pd.DataFrame()

        if df_ind_atual.empty:
            st.info("Nenhum ICP/Giro salvo para esta proposta. A DOA está usando valores médios padrão.")
        else:
            df_ind_atual["Explicação da nota"] = df_ind_atual.apply(
                lambda r: explicacao_nota_indicador(str(r.get("tipo", "")), str(r.get("indicador", "")), int(r.get("nota", 3) or 3)),
                axis=1
            )
            df_ind_visual = df_ind_atual.rename(columns={
                "tipo": "Tipo",
                "indicador": "Indicador",
                "nota": "Nota",
                "justificativa": "Justificativa"
            })[["Tipo", "Indicador", "Nota", "Explicação da nota", "Justificativa"]]
            st.dataframe(df_ind_visual, use_container_width=True, hide_index=True)

    if st.button("Salvar resultado da DOA", type="primary"):
        df_doa = carregar_doa()

        if not df_doa.empty:
            df_doa = df_doa[
                ~(
                    (df_doa["id_projeto"].astype(str) == str(id_atual)) &
                    (df_doa["versao"].astype(int) == int(versao_atual))
                )
            ]

        memoria_salvar = memoria.copy()
        memoria_salvar["id_projeto"] = id_atual
        memoria_salvar["versao"] = versao_atual
        memoria_salvar["valor_ajustado"] = memoria_salvar["valor_calculado"] * fator_limitador
        memoria_salvar["fator_limitador"] = fator_limitador
        memoria_salvar["data_calculo"] = agora()

        for coluna in colunas_doa():
            if coluna not in memoria_salvar.columns:
                memoria_salvar[coluna] = None

        memoria_salvar = memoria_salvar[colunas_doa()]
        df_doa = pd.concat([df_doa, memoria_salvar], ignore_index=True)
        salvar_excel(df_doa, ARQ_DOA)

        st.success("Resultado da DOA salvo com sucesso.")

# =========================
# PÁGINA 5 — PROPOSTA
# =========================

elif pagina == "5. Proposta Comercial":

    st.subheader("5. Proposta Comercial")

    if not id_atual:
        st.warning("Crie ou abra uma proposta antes de gerar o documento.")
        st.stop()

    st.markdown(
        "A proposta será gerada a partir do modelo `dados/modelo_proposta.docx`, "
        "incluindo os dados do projeto, o resumo das rubricas do PARF e o detalhamento final da DOA."
    )

    if not ARQ_MODELO_PROPOSTA.exists():
        st.error("Modelo não encontrado. Salve o arquivo como dados/modelo_proposta.docx.")
        st.stop()

    df_parf_atual = obter_parf_atual_df()
    if df_parf_atual.empty:
        st.warning("Nenhum item foi lançado no PARF desta proposta. Preencha o PARF antes de gerar a proposta.")
    else:
        st.markdown("### Prévia das informações que serão inseridas")

        resumo_parf = obter_resumo_parf_proposta(df_parf_atual)
        resumo_parf_visual = resumo_parf.copy()
        if not resumo_parf_visual.empty:
            resumo_parf_visual["Valor previsto"] = resumo_parf_visual["Valor previsto"].apply(moeda)
        st.markdown("#### Resumo das rubricas do PARF")
        st.dataframe(resumo_parf_visual, use_container_width=True, hide_index=True)

        resumo_doa = obter_resumo_doa_proposta(df_parf_atual)
        resumo_doa_visual = resumo_doa.copy()
        if not resumo_doa_visual.empty:
            resumo_doa_visual["Valor calculado"] = resumo_doa_visual["Valor calculado"].apply(moeda)
            resumo_doa_visual["Valor final"] = resumo_doa_visual["Valor final"].apply(moeda)
        st.markdown("#### Detalhamento da DOA")
        st.dataframe(resumo_doa_visual, use_container_width=True, hide_index=True)

    if st.button("Gerar proposta", type="primary"):
        caminho_proposta = gerar_proposta_docx()
        if caminho_proposta and caminho_proposta.exists():
            st.session_state["ultima_proposta_docx"] = str(caminho_proposta)
            st.success("Proposta gerada com sucesso.")

    caminho_ultima = st.session_state.get("ultima_proposta_docx")
    if caminho_ultima and Path(caminho_ultima).exists():
        with open(caminho_ultima, "rb") as f:
            st.download_button(
                label="Baixar proposta em Word",
                data=f,
                file_name=Path(caminho_ultima).name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

# =========================
# PÁGINA 6 — HISTÓRICO
# =========================

elif pagina == "6. Histórico":

    st.subheader("6. Histórico de Propostas")

    df = carregar_projetos()

    if df.empty:
        st.info("Nenhuma proposta cadastrada.")
    else:
        df_visual = df.rename(columns={
            "id_projeto": "ID da proposta",
            "versao": "Versão",
            "status": "Status",
            "nome_projeto": "Nome do projeto",
            "coordenador": "Coordenador",
            "email_coordenador": "E-mail do coordenador",
            "instituicao_executora": "Instituição executora",
            "campus": "Campus",
            "financiador": "Financiador",
            "tipo_financiador": "Tipo de financiador",
            "tipo_instrumento": "Tipo de instrumento",
            "data_inicio": "Data de início",
            "data_fim": "Data de fim",
            "prazo_meses": "Prazo (meses)",
            "valor_aprovado": "Valor aprovado/previsto",
            "objeto_resumido": "Objeto resumido",
            "area_tematica": "Área temática",
            "responsavel_facto": "Responsável Facto",
            "data_criacao": "Data de criação",
            "data_atualizacao": "Data de atualização"
        }).copy()
        if "Valor aprovado/previsto" in df_visual.columns:
            df_visual["Valor aprovado/previsto"] = df_visual["Valor aprovado/previsto"].apply(moeda)
        st.dataframe(df_visual, use_container_width=True, hide_index=True)